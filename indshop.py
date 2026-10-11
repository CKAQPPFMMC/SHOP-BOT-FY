# -*- coding: utf-8 -*-
"""
INDIA SHOP BOT  (INR / UPI auto payment via VC Gateway Pro)
Single file: bot.py  |  pip install -r requirements.txt  |  python bot.py

Setup: BOT_TOKEN aur OWNER_IDS neeche set karo (ya environment variable me do).
UPI ID aur Gateway API Key admin panel -> Payment Settings se bhi badal sakte ho.
"""
import asyncio
import base64
import hashlib
import hmac
import io
import json
import logging
import math
import os
import random
import re
import string
import struct
import time
from datetime import datetime, timedelta, timezone
from html import escape as esc

import aiohttp
import aiosqlite
from aiogram import BaseMiddleware, Bot, Dispatcher, F, Router
from aiogram.client.default import DefaultBotProperties
from aiogram.client.session.middlewares.base import BaseRequestMiddleware
from aiogram.enums import ParseMode
from aiogram.exceptions import TelegramBadRequest
from aiogram.filters import BaseFilter, Command, CommandStart
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.fsm.storage.memory import MemoryStorage
from aiogram.types import (BufferedInputFile, CallbackQuery, CopyTextButton, InlineKeyboardButton,
                           InlineKeyboardMarkup, KeyboardButton, Message, ReplyKeyboardMarkup)

# ======================================================================
#                              CONFIG
# ======================================================================
BOT_TOKEN = os.getenv("BOT_TOKEN", "PASTE_YOUR_BOT_TOKEN_HERE")
OWNER_IDS = [int(x) for x in os.getenv("OWNER_IDS", "123456789").split(",") if x.strip()]
SHOP_NAME = "Unlimited Shopping Bot"

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.getenv("DATA_DIR", os.path.join(BASE_DIR, "data"))
os.makedirs(DATA_DIR, exist_ok=True)
BACKUP_DIR = os.path.join(DATA_DIR, "backups")
os.makedirs(BACKUP_DIR, exist_ok=True)


def _find_db():
    for cand in (os.path.join(os.getcwd(), "shop.db"), os.path.join(BASE_DIR, "shop.db")):
        if os.path.exists(cand):
            return cand
    return os.path.join(DATA_DIR, "shop.db")


DB_PATH = os.getenv("DB_PATH") or _find_db()

MAX_QTY = 500                 # ek order me max pcs
MAX_DEPOSIT = 50000           # ek deposit me max Rs
MAX_BTN_ITEMS = 10            # itne accounts tak har account ka alag message + buttons, usse zyada ho to .txt file
BACKUP_EVERY_HOURS = 24
IST = timezone(timedelta(hours=5, minutes=30))
SHOW_FULL_NAME_IN_CHANNEL = False   # True karoge to channel notification me poora naam dikhega

# ---- VC Gateway Pro (UPI auto payment) ----
VC_API_URL = "https://vcgatewaypro.com/payment_api.php"
POLL_EVERY = 8                # sec: payment auto-check interval
POLL_MINUTES = 20             # itne minute tak auto check, phir 'expired'
CHECK_WINDOW_HOURS = 24       # expired payment ko bhi itne ghante tak 'Check Payment' se verify kar sakte hain

DEFAULT_SETTINGS = {
    "upi_id": os.getenv("UPI_ID", "paytm.s1dw5n0@pty"),
    "api_key": os.getenv("VC_API_KEY", "PAYB3495AF26D2371403DF25661"),
    "min_dep": "10", "support_user": "", "tutorial": "", "notify_chat": "",
}

# ======================================================================
#   BUTTON REGISTRY  key: (default emoji, label, color)   primary=blue success=green danger=red
# ======================================================================
B = {
    "buy": ("🛒", "Buy Products", "success"),
    "profile": ("👤", "My Profile", "primary"),
    "deposit": ("💰", "Deposit", "primary"),
    "getcode": ("🛡", "History", "primary"),
    "support": ("🎧", "Support", "primary"),
    "tutorial": ("👑", "Tutorial", "success"),
    "admin": ("🛠", "Admin Panel", "primary"),
    "a_prod": ("📦", "Product Management", "primary"),
    "a_addstock": ("🔥", "Add Stock", "success"),
    "a_live": ("📊", "Live Stock", "primary"),
    "a_users": ("👥", "User Management", "danger"),
    "a_depreq": ("🚨", "Deposit Requests", "success"),
    "a_dephist": ("🧾", "Deposit History", "primary"),
    "a_bcast": ("📢", "Broadcast", "danger"),
    "a_voucher": ("🎟", "Voucher Management", "success"),
    "a_force": ("🔒", "Force Join Setup", "primary"),
    "a_support": ("🎧", "Support Setup", "danger"),
    "a_admins": ("👮", "Admin Management", "primary"),
    "a_pay": ("💳", "Payment Settings", "success"),
    "a_resell": ("🤝", "Reseller Management", "primary"),
    "a_tutorial": ("📚", "Tutorial Setup", "success"),
    "a_emoji": ("✨", "Premium Emoji Set", "danger"),
    "a_stats": ("📈", "Statistics", "primary"),
    "a_channel": ("📣", "Channel Setup", "success"),
    "usermenu": ("✅", "User Menu", "success"),
    # inline
    "back": ("❌", "Back", "danger"),
    "chkpay": ("🔄", "Check Payment", "success"),
    "cancelpay": ("🚫", "Cancel", "danger"),
    "enterqty": ("✏️", "Enter Quantity", "success"),
    "applyvoucher": ("🎁", "Apply Voucher", "primary"),
    "buyhist": ("🛍", "Buy History", "primary"),
    "rechist": ("💵", "Recharge History", "success"),
    "track": ("🔎", "Track Order", "success"),
    "twofa": ("🔐", "2FA Code", "primary"),
    "contact": ("📞", "Contact Admin", "success"),
    "approve": ("✅", "Approve", "success"),
    "reject": ("❌", "Reject", "danger"),
    "ban": ("⛔", "Ban", "danger"),
    "unban": ("🟢", "Unban", "success"),
    "addbal": ("➕", "Add Balance", "primary"),
    "cutbal": ("➖", "Cut Balance", "danger"),
    "sendmsg": ("✉️", "Send Message", "primary"),
    "reply": ("↩️", "Reply", "success"),
    "addcat": ("📁", "Add Category", "success"),
    "delcat": ("🗑", "Delete Category", "danger"),
    "addprod": ("🆕", "Add Product", "primary"),
    "chprice": ("💲", "Change Price", "primary"),
    "delprod": ("❌", "Delete Product", "danger"),
    "delstock": ("🧹", "Delete Product Stock", "primary"),
    "addadmin": ("👑", "Add Admin", "success"),
    "deladmin": ("❌", "Delete Admin", "danger"),
    "vcreate": ("➕", "Create Voucher", "success"),
    "vdelete": ("🗑", "Delete Voucher", "danger"),
    "vlist": ("📋", "Voucher List", "primary"),
    "fadd": ("➕", "Add Channel", "success"),
    "fremove": ("🗑", "Remove Channel", "danger"),
    "addresell": ("➕", "Add Reseller", "success"),
    "delresell": ("➖", "Remove Reseller", "danger"),
    "yes": ("✅", "Yes, Confirm", "success"),
    "no": ("❌", "Cancel", "danger"),
    "joined": ("✅", "I Joined", "success"),
    "edit_upi": ("🏦", "Edit UPI ID", "success"),
    "edit_key": ("🔑", "Edit Gateway API Key", "primary"),
    "edit_min": ("⬇️", "Edit Minimum Deposit", "primary"),
}
MENU_KEYS = [k for k in B if k in (
    "buy", "profile", "deposit", "getcode", "support", "tutorial", "admin", "usermenu") or k.startswith("a_")]

# ======================================================================
#                         GLOBAL STATE / CACHE
# ======================================================================
DB: aiosqlite.Connection = None
bot: Bot = None
BOT_USERNAME = ""
SET: dict = {}
EMO: dict = {}
ADMINS: set = set()
TEMO: dict = {}
TEMO_RE = None
SUPMAP: dict = {}
PAYMSG: dict = {}        # deposit id -> (chat_id, message_id) payment screen
LAST_CHECK: dict = {}    # user id -> last 'Check Payment' time
TASKS: set = set()
BUY_LOCK = asyncio.Lock()
r = Router()


# ======================================================================
#                               DATABASE
# ======================================================================
async def init_db():
    global DB
    DB = await aiosqlite.connect(DB_PATH)
    DB.row_factory = aiosqlite.Row
    await DB.execute("PRAGMA journal_mode=WAL")
    await DB.execute("PRAGMA synchronous=NORMAL")
    await DB.execute("PRAGMA busy_timeout=10000")
    await DB.executescript("""
    CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY, name TEXT, username TEXT,
        balance REAL DEFAULT 0, banned INTEGER DEFAULT 0, reseller INTEGER DEFAULT 0,
        referrer INTEGER, joined INTEGER);
    CREATE TABLE IF NOT EXISTS admins(id INTEGER PRIMARY KEY);
    CREATE TABLE IF NOT EXISTS categories(id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT, emoji_id TEXT);
    CREATE TABLE IF NOT EXISTS products(id INTEGER PRIMARY KEY AUTOINCREMENT, cat_id INTEGER, name TEXT,
        price REAL, rprice REAL, emoji_id TEXT);
    CREATE TABLE IF NOT EXISTS stock(id INTEGER PRIMARY KEY AUTOINCREMENT, product_id INTEGER, data TEXT,
        sold INTEGER DEFAULT 0, buyer INTEGER, sold_at INTEGER);
    CREATE INDEX IF NOT EXISTS ix_stock ON stock(product_id, sold);
    CREATE TABLE IF NOT EXISTS orders(id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER, product_id INTEGER,
        product_name TEXT, cat_name TEXT, invoice TEXT, qty INTEGER, total REAL, created INTEGER, items TEXT);
    CREATE TABLE IF NOT EXISTS deposits(id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER, method TEXT,
        amount REAL, trx TEXT, photo TEXT, status TEXT DEFAULT 'pending', created INTEGER, decided INTEGER);
    CREATE TABLE IF NOT EXISTS settings(key TEXT PRIMARY KEY, value TEXT);
    CREATE TABLE IF NOT EXISTS emojis(key TEXT PRIMARY KEY, emoji_id TEXT);
    CREATE TABLE IF NOT EXISTS text_emojis(ch TEXT PRIMARY KEY, emoji_id TEXT);
    CREATE TABLE IF NOT EXISTS vouchers(code TEXT PRIMARY KEY, amount REAL, max_uses INTEGER, used INTEGER DEFAULT 0);
    CREATE TABLE IF NOT EXISTS voucher_uses(code TEXT, user_id INTEGER);
    CREATE TABLE IF NOT EXISTS force_channels(id INTEGER PRIMARY KEY AUTOINCREMENT, chat TEXT, title TEXT, link TEXT);
    """)
    for stmt in ("ALTER TABLE orders ADD COLUMN cat_name TEXT", "ALTER TABLE orders ADD COLUMN invoice TEXT"):
        try:
            await DB.execute(stmt)
        except Exception:
            pass
    for k, v in DEFAULT_SETTINGS.items():
        await DB.execute("INSERT OR IGNORE INTO settings(key,value) VALUES(?,?)", (k, v))
    await DB.commit()
    await load_cache()


async def load_cache():
    SET.clear(); EMO.clear(); ADMINS.clear(); TEMO.clear()
    for x in await fetchall("SELECT * FROM settings"):
        SET[x["key"]] = x["value"]
    for x in await fetchall("SELECT * FROM emojis"):
        EMO[x["key"]] = x["emoji_id"]
    for x in await fetchall("SELECT id FROM admins"):
        ADMINS.add(x["id"])
    for x in await fetchall("SELECT * FROM text_emojis"):
        TEMO[x["ch"]] = x["emoji_id"]
    rebuild_temo()


async def fetchone(sql, args=()):
    async with DB.execute(sql, args) as c:
        return await c.fetchone()


async def fetchall(sql, args=()):
    async with DB.execute(sql, args) as c:
        return await c.fetchall()


async def run(sql, args=()):
    cur = await DB.execute(sql, args)
    await DB.commit()
    return cur


async def set_setting(k, v):
    SET[k] = str(v)
    await run("INSERT OR REPLACE INTO settings(key,value) VALUES(?,?)", (k, str(v)))


# ======================================================================
#                               HELPERS
# ======================================================================
VS16 = "\ufe0f"
_TAG_RE = re.compile(r"(<tg-emoji[^>]*>.*?</tg-emoji>|<[^>]+>)", re.S)
PE_METHODS = {"SendMessage", "EditMessageText", "SendPhoto", "SendDocument", "SendVideo",
              "EditMessageCaption", "CopyMessage"}


def _nk(ch):
    return ch.replace(VS16, "")


def rebuild_temo():
    global TEMO_RE
    if TEMO:
        keys = sorted(TEMO, key=len, reverse=True)
        TEMO_RE = re.compile("(" + "|".join(re.escape(k) for k in keys) + ")" + VS16 + "?")
    else:
        TEMO_RE = None


def pe(text):
    """Text ke normal emoji ko premium emoji se replace karta hai (Text Emoji setting)"""
    if not TEMO_RE or not text:
        return text
    out, code = [], 0
    for i, part in enumerate(_TAG_RE.split(text)):
        if i % 2 == 1:
            low = part.lower()
            if low in ("<code>", "<pre>"):
                code += 1
            elif low in ("</code>", "</pre>"):
                code = max(0, code - 1)
            out.append(part)
        elif code:
            out.append(part)
        else:
            out.append(TEMO_RE.sub(lambda mo: f'<tg-emoji emoji-id="{TEMO[mo.group(1)]}">{mo.group(0)}</tg-emoji>', part))
    return "".join(out)


class PremiumText(BaseRequestMiddleware):
    async def __call__(self, make_request, bot, method):
        if TEMO_RE and type(method).__name__ in PE_METHODS:
            pm = getattr(method, "parse_mode", None)
            if not isinstance(pm, str) or str(pm).lower().endswith("html"):
                upd = {}
                for f in ("text", "caption"):
                    v = getattr(method, f, None)
                    if isinstance(v, str):
                        upd[f] = pe(v)
                if upd:
                    method = method.model_copy(update=upd)
        return await make_request(bot, method)


def parse_amount(t, maxv):
    try:
        v = float(t.strip().replace(",", ""))
    except ValueError:
        return None
    if not math.isfinite(v) or v <= 0 or v > maxv:
        return None
    return v


def admin_ids():
    return set(OWNER_IDS) | ADMINS


def is_admin(uid):
    return uid in OWNER_IDS or uid in ADMINS


def money(x):
    return f"₹{x:.2f}"


def fp(x):
    return f"{x:.1f}"


def now():
    return int(time.time())


def day_bounds():
    t0 = datetime.now(IST).replace(hour=0, minute=0, second=0, microsecond=0)
    return int(t0.timestamp()), int((t0 - timedelta(days=1)).timestamp())


def fdate(ts):
    return datetime.fromtimestamp(ts, IST).strftime("%d-%b %I:%M %p")


def em(key):
    fb = B[key][0] if key in B else "•"
    cid = EMO.get(key)
    return f'<tg-emoji emoji-id="{cid}">{fb}</tg-emoji>' if cid else fb


def cem(cid, fb):
    return f'<tg-emoji emoji-id="{cid}">{fb}</tg-emoji>' if cid else fb


def IM(rows):
    return InlineKeyboardMarkup(inline_keyboard=rows)


def ib(key, text=None, **kw):
    emo, label, style = B[key]
    cid = EMO.get(key)
    t = text or label
    d = dict(text=t if cid else f"{emo} {t}", style=style, **kw)
    if cid:
        d["icon_custom_emoji_id"] = cid
    return InlineKeyboardButton(**d)


def dib(text, cb, cid=None, default="📁", style="primary"):
    d = dict(text=text if cid else f"{default} {text}", callback_data=cb, style=style)
    if cid:
        d["icon_custom_emoji_id"] = cid
    return InlineKeyboardButton(**d)


def copy_btn(text, value, style="primary"):
    return InlineKeyboardButton(text=text, copy_text=CopyTextButton(text=value[:256]), style=style)


def rb(key):
    emo, label, style = B[key]
    cid = EMO.get(key)
    d = dict(text=label if cid else f"{emo} {label}", style=style)
    if cid:
        d["icon_custom_emoji_id"] = cid
    return KeyboardButton(**d)


def RK(rows):
    return ReplyKeyboardMarkup(keyboard=rows, resize_keyboard=True)


def main_kb(uid):
    rows = [[rb("buy")], [rb("profile"), rb("deposit")], [rb("getcode"), rb("support")], [rb("tutorial")]]
    if is_admin(uid):
        rows.append([rb("admin")])
    return RK(rows)


def admin_kb():
    k = ["a_prod", "a_addstock", "a_live", "a_users", "a_depreq", "a_dephist", "a_bcast", "a_voucher",
         "a_force", "a_support", "a_admins", "a_pay", "a_resell", "a_tutorial", "a_emoji", "a_stats", "a_channel"]
    rows = [[rb(k[i])] + ([rb(k[i + 1])] if i + 1 < len(k) else []) for i in range(0, len(k), 2)]
    rows.append([rb("usermenu")])
    return RK(rows)


async def show(ev, text, kb=None):
    if isinstance(ev, CallbackQuery):
        try:
            await ev.message.edit_text(text, reply_markup=kb)
            return
        except TelegramBadRequest:
            await ev.message.answer(text, reply_markup=kb)
            return
    await ev.answer(text, reply_markup=kb)


async def notify_admins(text, **kw):
    for a in admin_ids():
        try:
            await bot.send_message(a, text, **kw)
        except Exception:
            pass


async def notify_channel(text):
    ch = (SET.get("notify_chat") or "").strip()
    if not ch:
        return
    try:
        await bot.send_message(int(ch) if ch.lstrip("-").isdigit() else ch, text)
    except Exception:
        logging.exception("channel notify failed")


async def ensure_user(u, ref=None):
    row = await fetchone("SELECT id FROM users WHERE id=?", (u.id,))
    if not row:
        await DB.execute("INSERT INTO users(id,name,username,joined,referrer) VALUES(?,?,?,?,?)",
                         (u.id, u.full_name, u.username or "", now(), ref))
    else:
        await DB.execute("UPDATE users SET name=?, username=? WHERE id=?", (u.full_name, u.username or "", u.id))
    await DB.commit()


async def get_user(uid):
    return await fetchone("SELECT * FROM users WHERE id=?", (uid,))


async def find_user(txt):
    txt = txt.strip().lstrip("@")
    if txt.isdigit():
        return await fetchone("SELECT * FROM users WHERE id=?", (int(txt),))
    return await fetchone("SELECT * FROM users WHERE lower(username)=lower(?)", (txt,))


async def stock_count(pid):
    x = await fetchone("SELECT COUNT(*) c FROM stock WHERE product_id=? AND sold=0", (pid,))
    return x["c"]


def totp(secret):
    s = re.sub(r"\s+", "", secret).upper()
    s += "=" * (-len(s) % 8)
    key = base64.b32decode(s, casefold=True)
    h = hmac.new(key, struct.pack(">Q", int(time.time() // 30)), hashlib.sha1).digest()
    o = h[-1] & 15
    code = (struct.unpack(">I", h[o:o + 4])[0] & 0x7FFFFFFF) % 1000000
    return f"{code:06d}", 30 - int(time.time()) % 30


def parse_emoji_id(m: Message):
    for e in (m.entities or []):
        if e.type == "custom_emoji":
            return e.custom_emoji_id
    t = (m.text or "").strip()
    return t if t.isdigit() else None


# ---------------- stock parsing: Link | Username | Password ----------------
URL_RE = re.compile(r"https?://[^\s)\]>\"']+")
HEAD_RE = re.compile(r"^\s*(login link|login url|link|url)\b", re.I)


def norm_line(line):
    line = line.strip()
    if "|" in line:
        parts = [x.strip() for x in line.split("|", 2)]
    elif "\t" in line:
        parts = [x.strip() for x in line.split("\t", 2)]
    elif line.count(",") >= 2:
        parts = [x.strip() for x in line.split(",", 2)]
    else:
        parts = line.split(None, 2)
    if len(parts) == 3 and all(parts):
        return "|".join(parts)
    return line


def _cell_str(c):
    if isinstance(c, float) and c.is_integer():
        c = int(c)
    return str(c).strip()


def build_items(rows):
    """Layouts: A=Link,B=Username,C=Password | 'Login Link (https://..)' header + A=Username,B=Password | ek cell 'link | user | pass'"""
    out, default = [], None
    for cells in rows:
        if not cells:
            continue
        first = cells[0]
        if len(cells) == 1 and HEAD_RE.match(first):
            m = URL_RE.search(first)
            if m:
                default = m.group(0)
            continue
        if first.lower() in ("login link", "login url", "link", "url"):
            continue
        if len(cells) >= 3:
            out.append("|".join(cells[:3]))
        elif len(cells) == 2 and default:
            out.append(f"{default}|{cells[0]}|{cells[1]}")
        elif len(cells) == 1:
            out.append(norm_line(first))
        else:
            out.append(" ".join(cells))
    return out


def parse_txt(b):
    rows = []
    for line in b.decode("utf-8", "ignore").splitlines():
        line = line.strip()
        if not line:
            continue
        if "|" not in line and HEAD_RE.match(line) and URL_RE.search(line):
            cells = [line]
        elif "|" in line:
            cells = [x.strip() for x in line.split("|", 2)]
        elif "\t" in line:
            cells = [x.strip() for x in line.split("\t", 2)]
        elif line.count(",") >= 2:
            cells = [x.strip() for x in line.split(",", 2)]
        else:
            cells = line.split(None, 2)
        rows.append([c for c in cells if c])
    return build_items(rows)


def parse_xlsx(b):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(b), read_only=True, data_only=True)
    rows = []
    for row in wb.worksheets[0].iter_rows(values_only=True):
        rows.append([_cell_str(c) for c in row if c is not None and _cell_str(c) != ""])
    return build_items(rows)


def parse_xls(b):
    import xlrd
    sh = xlrd.open_workbook(file_contents=b).sheet_by_index(0)
    rows = []
    for i in range(sh.nrows):
        rows.append([_cell_str(c) for c in sh.row_values(i) if c not in (None, "") and _cell_str(c) != ""])
    return build_items(rows)


def parse_stock_file(name, data):
    name = (name or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        return parse_xlsx(data)
    if name.endswith(".xls"):
        return parse_xls(data)
    return parse_txt(data)


# ---------------- delivery format ----------------
def fmt_item(data):
    p = data.split("|", 2)
    if len(p) == 3:
        return (f"🔗 <b>Login Link:</b> <code>{esc(p[0])}</code>\n👤 <b>Username:</b> <code>{esc(p[1])}</code>\n"
                f"🔑 <b>Password:</b> <code>{esc(p[2])}</code>")
    return f"<code>{esc(data)}</code>"


def norm_url(u):
    u = u.strip()
    if not re.match(r"^https?://", u, re.I):
        u = "https://" + u.lstrip("/")
    return u


def _short(s, n=28):
    return s if len(s) <= n else s[:n - 1] + "…"


def item_kb(data):
    """Sabse upar Login Link (open), neeche left Username (copy) | right Password (copy)"""
    p = data.split("|", 2)
    if len(p) != 3:
        return None
    link, user, pw = (x.strip() for x in p)
    rows = []
    url = norm_url(link)
    if re.match(r"^https?://[^\s/]+\.[^\s/]+\S*$", url) and len(url) < 500:
        rows.append([InlineKeyboardButton(text="🔗 Open Login Link", url=url, style="success")])
    copy_row = []
    if user and len(user) <= 256:
        copy_row.append(copy_btn(f"👤 {_short(user)}", user, "primary"))
    if pw and len(pw) <= 256:
        copy_row.append(copy_btn(f"🔑 {_short(pw)}", pw, "danger"))
    if copy_row:
        rows.append(copy_row)
    return IM(rows) if rows else None


async def send_items(target, items, with_buttons=True):
    n = len(items)
    if n > MAX_BTN_ITEMS:
        txt = "\n".join(x.replace("|", " | ") for x in items)
        await target.answer_document(BufferedInputFile(txt.encode(), f"accounts_{now()}.txt"),
                                     caption="📄 Saare accounts is file me hain (Link | Username | Password)")
        return
    for i, x in enumerate(items, 1):
        head = f"📦 <b>Account {i}/{n}</b>\n" if n > 1 else ""
        txt = head + fmt_item(x) + "\n\n👇 Link kholne ke liye aur Username/Password copy karne ke liye buttons dabao"
        kb = item_kb(x) if with_buttons else None
        try:
            await target.answer(txt, reply_markup=kb)
        except TelegramBadRequest:
            await target.answer(head + fmt_item(x))


# ---------------- masks for channel ----------------
def mask_name(n):
    if SHOW_FULL_NAME_IN_CHANNEL:
        return n
    return " ".join(w[:2] + "***" for w in (n or "").split()) or "User"


def mask_user(u):
    if not u:
        return "Hidden"
    return "@" + (u[:2] + "***" + u[-2:] if len(u) > 5 else u[:1] + "***")


def mask_id(i):
    s = str(i)
    return s[:3] + "***" + s[-2:]


# ---------------- UPI QR + gateway ----------------
def make_qr_png(data):
    try:
        import qrcode
        img = qrcode.make(data)
        bio = io.BytesIO()
        img.save(bio, "PNG")
        return bio.getvalue()
    except Exception:
        return None


KNOWN = ("NOT_FOUND", "INVALID", "FAILED", "PENDING", "SUCCESS")
SYN = {"SUCCESSFUL": "SUCCESS", "PAID": "SUCCESS", "COMPLETED": "SUCCESS", "COMPLETE": "SUCCESS",
       "FAIL": "FAILED", "FAILURE": "FAILED", "NOTFOUND": "NOT_FOUND", "NOT FOUND": "NOT_FOUND"}


def _norm_status(v):
    s = str(v).strip().upper().replace("-", "_")
    s = SYN.get(s, s)
    return s if s in KNOWN else None


def parse_vc(raw, amount):
    """Gateway response -> SUCCESS / PENDING / FAILED / INVALID / NOT_FOUND / AMOUNT_MISMATCH / UNKNOWN"""
    try:
        data = json.loads(raw)
    except Exception:
        data = None
    if isinstance(data, dict):
        flat = dict(data)
        for k in ("data", "result", "payment"):
            if isinstance(data.get(k), dict):
                flat.update(data[k])
        low = {str(k).lower(): v for k, v in flat.items()}
        st = None
        for k in ("status", "payment_status", "state", "result"):
            if k in low and not isinstance(low[k], (dict, list)):
                st = _norm_status(low[k])
                if st:
                    break
        if st == "SUCCESS":
            for k in ("amount", "paid_amount", "txn_amount"):
                if k in low:
                    try:
                        if abs(float(str(low[k]).replace(",", "")) - float(amount)) > 0.01:
                            return "AMOUNT_MISMATCH"
                    except ValueError:
                        pass
                    break
        return st or "UNKNOWN"
    up = (raw or "").strip().upper()
    st = _norm_status(up)
    if st:
        return st
    for k in ("NOT_FOUND", "INVALID", "FAILED", "PENDING"):
        if re.search(r"\b" + k + r"\b", up):
            return k
    if re.search(r"\bSUCCESS\b", up):
        return "SUCCESS"
    return "UNKNOWN"


async def vc_call(order_id, amount):
    params = {"api_key": SET.get("api_key", ""), "order_id": order_id, "amount": str(int(amount))}
    async with aiohttp.ClientSession(timeout=aiohttp.ClientTimeout(total=25)) as s:
        async with s.get(VC_API_URL, params=params) as resp:
            raw = await resp.text()
    logging.info("VC verify %s -> %s", order_id, raw[:200].replace("\n", " "))
    return parse_vc(raw, amount), raw


async def new_order_id():
    while True:
        rnd = "".join(random.choice(string.digits + string.ascii_uppercase) for _ in range(6))
        oid = "ORD" + datetime.now(IST).strftime("%y%m%d") + rnd
        if not await fetchone("SELECT 1 FROM deposits WHERE trx=?", (oid,)):
            return oid


# ======================================================================
#                                STATES
# ======================================================================
class S(StatesGroup):
    dep_amount = State(); qty = State(); support = State(); voucher = State(); twofa = State(); track = State()
    a_reply = State()
    a_cat_name = State(); a_prod_name = State(); a_prod_price = State(); a_prod_rprice = State()
    a_price = State(); a_rprice = State(); a_stock = State()
    a_user = State(); a_bal = State(); a_msg = State(); a_bcast = State()
    a_vcode = State(); a_vamt = State(); a_vuses = State()
    a_temoji = State(); a_temoji_del = State()
    a_force = State(); a_admin = State(); a_resell = State(); a_setval = State(); a_emoji = State()
    a_channel = State()


# ======================================================================
#                      FILTERS + MIDDLEWARE
# ======================================================================
class AdminF(BaseFilter):
    async def __call__(self, ev):
        return is_admin(ev.from_user.id)


A = AdminF()


class MenuF(BaseFilter):
    async def __call__(self, m: Message):
        t = (m.text or "").strip()
        if not t:
            return False
        for k in MENU_KEYS:
            label = B[k][1]
            if t.endswith(label):
                prefix = t[:-len(label)].strip()
                if not prefix or (len(prefix) <= 8 and not any(ch.isalnum() for ch in prefix)):
                    return {"mkey": k}
        return False


class SupReplyF(BaseFilter):
    async def __call__(self, m: Message):
        if not m.reply_to_message or not is_admin(m.from_user.id):
            return False
        uid = SUPMAP.get((m.chat.id, m.reply_to_message.message_id))
        return {"sup_uid": uid} if uid else False


async def missing_channels(uid):
    miss = []
    for c in await fetchall("SELECT * FROM force_channels"):
        try:
            chat = int(c["chat"]) if c["chat"].lstrip("-").isdigit() else c["chat"]
            mem = await bot.get_chat_member(chat, uid)
            if mem.status in ("left", "kicked"):
                miss.append(c)
        except Exception:
            pass
    return miss


class Guard(BaseMiddleware):
    async def __call__(self, handler, event, data):
        user = data.get("event_from_user")
        if not user or user.is_bot:
            return await handler(event, data)
        ref = None
        if isinstance(event, Message) and event.text and event.text.startswith("/start "):
            arg = event.text.split(maxsplit=1)[1].strip()
            if arg.isdigit() and int(arg) != user.id:
                ref = int(arg)
        await ensure_user(user, ref)
        if is_admin(user.id):
            return await handler(event, data)
        u = await get_user(user.id)
        if u["banned"]:
            if isinstance(event, CallbackQuery):
                await event.answer("🚫 Aapko ban kar diya gaya hai.", show_alert=True)
            else:
                await event.answer("🚫 <b>Aapko ban kar diya gaya hai.</b>")
            return
        if not (isinstance(event, CallbackQuery) and event.data == "chk"):
            miss = await missing_channels(user.id)
            if miss:
                rows = [[InlineKeyboardButton(text=f"📢 Join {c['title'] or c['chat']}", url=c["link"], style="primary")]
                        for c in miss if c["link"]]
                rows.append([ib("joined", callback_data="chk")])
                tgt = event.message if isinstance(event, CallbackQuery) else event
                await tgt.answer("🔒 <b>Bot use karne ke liye pehle neeche ke channels join karo:</b>", reply_markup=IM(rows))
                if isinstance(event, CallbackQuery):
                    await event.answer()
                return
        return await handler(event, data)


# ======================================================================
#   MENU DISPATCH - sabse pehle register (kisi bhi state me menu button kaam kare)
# ======================================================================
MENU_H = {}


@r.message(MenuF())
async def on_menu(m: Message, state: FSMContext, mkey: str):
    await state.clear()
    if (mkey.startswith("a_") or mkey in ("admin", "usermenu")) and not is_admin(m.from_user.id):
        return
    await MENU_H[mkey](m, state)


def welcome_text(name):
    return (f"{em('tutorial')} <b>Welcome to {esc(SHOP_NAME)}</b>\n\n"
            f"👋 <b>{esc(name)}</b>, neeche menu se option chuno.")


@r.message(CommandStart())
async def cmd_start(m: Message, state: FSMContext):
    await state.clear()
    await m.answer(welcome_text(m.from_user.full_name), reply_markup=main_kb(m.from_user.id))


@r.message(Command("cancel"))
async def cmd_cancel(m: Message, state: FSMContext):
    await state.clear()
    await m.answer("✅ <b>Cancelled</b>", reply_markup=main_kb(m.from_user.id))


@r.callback_query(F.data == "chk")
async def cb_chk(c: CallbackQuery):
    if await missing_channels(c.from_user.id):
        await c.answer("❌ Pehle saare channels join karo.", show_alert=True)
        return
    await c.answer()
    await c.message.answer(welcome_text(c.from_user.full_name), reply_markup=main_kb(c.from_user.id))


@r.callback_query(F.data == "noop")
async def cb_noop(c: CallbackQuery):
    await c.answer()


# ======================================================================
#                              BUY FLOW
# ======================================================================
async def cats_view():
    cats = await fetchall("SELECT * FROM categories ORDER BY id")
    if not cats:
        return "❌ <b>Abhi koi category nahi hai.</b>", None
    rows = []
    for i in range(0, len(cats), 2):
        rows.append([dib(c["name"], f"cat:{c['id']}", c["emoji_id"]) for c in cats[i:i + 2]])
    return f"{em('buy')} <b>Category select karo</b>", IM(rows)


async def h_buy(m, state):
    t, kb = await cats_view()
    await m.answer(t, reply_markup=kb)


async def prods_view(cid, uid):
    cat = await fetchone("SELECT * FROM categories WHERE id=?", (cid,))
    if not cat:
        return "❌ Category nahi mili.", None
    u = await get_user(uid)
    prods = await fetchall("SELECT * FROM products WHERE cat_id=? ORDER BY id", (cid,))
    rows = []
    for p in prods:
        price = p["rprice"] if u["reseller"] else p["price"]
        s = await stock_count(p["id"])
        rows.append([dib(f"{p['name']} | ₹{price:g} | Stock: {s}", f"prod:{p['id']}", p["emoji_id"], "🌐")])
    rows.append([ib("back", callback_data="buy_back")])
    return f"{cem(cat['emoji_id'], '🌐')} <b>{esc(cat['name'])} — Package select karo:</b>", IM(rows)


@r.callback_query(F.data == "buy_back")
async def cb_buy_back(c: CallbackQuery):
    await c.answer()
    t, kb = await cats_view()
    await show(c, t, kb)


@r.callback_query(F.data.startswith("cat:"))
async def cb_cat(c: CallbackQuery):
    await c.answer()
    t, kb = await prods_view(int(c.data.split(":")[1]), c.from_user.id)
    await show(c, t, kb)


@r.callback_query(F.data.startswith("prod:"))
async def cb_prod(c: CallbackQuery):
    await c.answer()
    pid = int(c.data.split(":")[1])
    p = await fetchone("SELECT * FROM products WHERE id=?", (pid,))
    if not p:
        return
    u = await get_user(c.from_user.id)
    price = p["rprice"] if u["reseller"] else p["price"]
    s = await stock_count(pid)
    txt = (f"{em('tutorial')} <b>Kitne lene hain?</b>\n━━━━━━━━━━━━━━\n"
           f"📌 Product: <b>{esc(p['name'])}</b>\n💲 Price: <b>₹{price:g} / pcs</b>\n🔥 Stock: <b>{s} pcs</b>")
    kb = IM([
        [dib("1 pcs", f"q:{pid}:1", default="👑"), dib("3 pcs", f"q:{pid}:3", default="👑")],
        [dib("5 pcs", f"q:{pid}:5", default="👑"), dib("10 pcs", f"q:{pid}:10", default="👑", style="danger")],
        [ib("enterqty", callback_data=f"qe:{pid}")],
        [ib("back", callback_data=f"cat:{p['cat_id']}")],
    ])
    await show(c, txt, kb)


async def purchase(uid, pid, n):
    async with BUY_LOCK:
        u = await get_user(uid)
        p = await fetchone("SELECT * FROM products WHERE id=?", (pid,))
        if not p:
            return "err", "❌ Product nahi mila."
        if n < 1 or n > MAX_QTY:
            return "err", f"❌ Minimum 1 aur maximum {MAX_QTY} pcs le sakte ho."
        rows = await fetchall("SELECT id,data FROM stock WHERE product_id=? AND sold=0 ORDER BY id LIMIT ?", (pid, n))
        if len(rows) < n:
            return "err", f"❌ Itna stock nahi hai. Available: {len(rows)} pcs"
        price = p["rprice"] if u["reseller"] else p["price"]
        total = round(price * n, 2)
        if u["balance"] + 1e-9 < total:
            return "nobal", total
        cat = await fetchone("SELECT name FROM categories WHERE id=?", (p["cat_id"],))
        ids = [x["id"] for x in rows]
        items = [x["data"] for x in rows]
        invoice = f"INV{int(time.time() * 1000)}"
        while await fetchone("SELECT 1 FROM orders WHERE invoice=?", (invoice,)):
            invoice = f"INV{int(time.time() * 1000) + random.randint(1, 999)}"
        try:
            ph = ",".join("?" * len(ids))
            await DB.execute(f"UPDATE stock SET sold=1, buyer=?, sold_at=? WHERE id IN ({ph})", [uid, now(), *ids])
            await DB.execute("UPDATE users SET balance=balance-? WHERE id=?", (total, uid))
            cur = await DB.execute(
                "INSERT INTO orders(user_id,product_id,product_name,cat_name,invoice,qty,total,created,items) VALUES(?,?,?,?,?,?,?,?,?)",
                (uid, pid, p["name"], cat["name"] if cat else "-", invoice, n, total, now(), "\n".join(items)))
            await DB.commit()
        except Exception:
            await DB.rollback()
            raise
        left = await stock_count(pid)
        u2 = await get_user(uid)
        return "ok", dict(oid=cur.lastrowid, invoice=invoice, items=items, total=total, name=p["name"],
                          cat=cat["name"] if cat else "-", left=left, bal=u2["balance"])


async def do_buy(target: Message, uid, pid, n):
    st, res = await purchase(uid, pid, n)
    if st == "err":
        await target.answer(res)
        return
    if st == "nobal":
        await target.answer(f"❌ <b>Balance kam hai!</b>\nChahiye: <b>{money(res)}</b>\nPehle <b>Deposit</b> karo.")
        return
    await target.answer(
        f"✅ <b>Purchase Successful!</b>\n━━━━━━━━━━━━━━\n🧾 Invoice: <code>{res['invoice']}</code>\n"
        f"📁 Category: <b>{esc(res['cat'])}</b>\n📌 Package: <b>{esc(res['name'])}</b>\n🔢 Quantity: <b>{n} pcs</b>\n"
        f"💰 Cost: <b>{money(res['total'])}</b>\n💳 Balance: <b>{money(res['bal'])}</b>\n\n"
        f"ℹ️ Order track karne ke liye: <b>History → Track Order</b> me Invoice ID daalo.")
    await send_items(target, res["items"])
    u = await get_user(uid)
    await notify_admins(
        f"{em('buy')} <b>New Purchase!</b>\n━━━━━━━━━━━━━━\n👤 {esc(u['name'])} (@{esc(u['username'] or '-')})\n"
        f"🆔 <code>{uid}</code>\n🧾 <code>{res['invoice']}</code>\n📁 {esc(res['cat'])} → {esc(res['name'])}\n"
        f"🔢 Qty: <b>{n}</b>\n💰 Total: <b>{money(res['total'])}</b>\n📦 Stock left: <b>{res['left']}</b>")
    await notify_channel(channel_text(res["invoice"], u["name"], u["username"], uid, res["cat"], res["name"], n,
                                      res["total"], now()))


def channel_text(inv, name, username, uid, cat, pkg, n, total, ts):
    return (f"✅ <b>AUTO DELIVERY SUCCESS!</b>\n\n🧾 Invoice: <code>{inv}</code>\n👤 Name: {esc(mask_name(name))}\n"
            f"🔰 Username: {esc(mask_user(username))}\n🆔 User ID: {mask_id(uid)}\n\n🛒 Category: {esc(cat)}\n"
            f"📁 Package: {esc(pkg)}\n🔢 Quantity: {n}\n💰 Amount: ₹{fp(total)}\n"
            f"📅 Time: {datetime.fromtimestamp(ts, IST).strftime('%Y-%m-%d %I:%M %p')}")


@r.callback_query(F.data.startswith("q:"))
async def cb_qty(c: CallbackQuery):
    await c.answer()
    _, pid, n = c.data.split(":")
    await do_buy(c.message, c.from_user.id, int(pid), int(n))


@r.callback_query(F.data.startswith("qe:"))
async def cb_qty_enter(c: CallbackQuery, state: FSMContext):
    await c.answer()
    await state.set_state(S.qty)
    await state.update_data(pid=int(c.data.split(":")[1]))
    await c.message.answer("✏️ <b>Kitne pcs chahiye? Number likho:</b>")


@r.message(S.qty, F.text)
async def st_qty(m: Message, state: FSMContext):
    if not m.text.strip().isdigit():
        await m.answer("❌ Sirf number likho.")
        return
    d = await state.get_data()
    await state.clear()
    await do_buy(m, m.from_user.id, d["pid"], int(m.text.strip()))


# ======================================================================
#                          PROFILE / VOUCHER
# ======================================================================
async def h_profile(m, state):
    u = await get_user(m.from_user.id)
    o = await fetchone("SELECT COALESCE(SUM(qty),0) q FROM orders WHERE user_id=?", (u["id"],))
    rf = await fetchone("SELECT COUNT(*) c FROM users WHERE referrer=?", (u["id"],))
    reseller_line = "🤝 Account: <b>Reseller</b>\n" if u["reseller"] else ""
    txt = (f"{em('profile')} <b>My Profile</b>\n━━━━━━━━━━━━━━\n"
           f"👑 Name: <b>{esc(u['name'])}</b>\n⚙️ User ID: <code>{u['id']}</code>\n"
           f"🌐 Username: @{esc(u['username'] or '-')}\n💰 Balance: <b>{money(u['balance'])}</b>\n"
           f"🧾 Total Orders: <b>{o['q']}</b>\n👥 Referrals: <b>{rf['c']}</b>\n"
           f"{reseller_line}━━━━━━━━━━━━━━\n"
           f"🔗 Referral Link:\n<code>https://t.me/{BOT_USERNAME}?start={u['id']}</code>")
    await m.answer(txt, reply_markup=IM([[ib("applyvoucher", callback_data="vch")]]))


@r.callback_query(F.data == "vch")
async def cb_vch(c: CallbackQuery, state: FSMContext):
    await c.answer()
    await state.set_state(S.voucher)
    await c.message.answer("🎁 <b>Apna Voucher Code likho:</b>")


@r.message(S.voucher, F.text)
async def st_voucher(m: Message, state: FSMContext):
    await state.clear()
    code = m.text.strip().upper()
    v = await fetchone("SELECT * FROM vouchers WHERE code=?", (code,))
    if not v:
        await m.answer("❌ <b>Voucher code galat hai.</b>")
        return
    if v["used"] >= v["max_uses"]:
        await m.answer("❌ <b>Is voucher ki limit khatam ho gayi.</b>")
        return
    if await fetchone("SELECT 1 FROM voucher_uses WHERE code=? AND user_id=?", (code, m.from_user.id)):
        await m.answer("❌ <b>Aap ye voucher pehle use kar chuke ho.</b>")
        return
    await DB.execute("UPDATE vouchers SET used=used+1 WHERE code=?", (code,))
    await DB.execute("INSERT INTO voucher_uses(code,user_id) VALUES(?,?)", (code, m.from_user.id))
    await DB.execute("UPDATE users SET balance=balance+? WHERE id=?", (v["amount"], m.from_user.id))
    await DB.commit()
    await m.answer(f"✅ <b>Voucher Applied!</b>\n+{money(v['amount'])} balance me add ho gaya.")


# ======================================================================
#                      DEPOSIT (UPI auto payment)
# ======================================================================
async def h_deposit(m, state):
    mn = float(SET.get("min_dep") or 10)
    await state.set_state(S.dep_amount)
    await m.answer(f"{em('deposit')} <b>UPI Deposit</b>\n━━━━━━━━━━━━━━\nKitne ₹ deposit karne hain?\n"
                   f"Minimum: <b>₹{mn:g}</b> | Maximum: <b>₹{MAX_DEPOSIT}</b>\nSirf number likho, jaise: <code>100</code>")


@r.message(S.dep_amount, F.text)
async def st_dep_amount(m: Message, state: FSMContext):
    t = m.text.strip()
    mn = float(SET.get("min_dep") or 10)
    if not t.isdigit() or int(t) <= 0:
        await m.answer("❌ Sirf poora number likho (jaise 100).")
        return
    amount = int(t)
    if amount < mn or amount > MAX_DEPOSIT:
        await m.answer(f"❌ Amount ₹{mn:g} se ₹{MAX_DEPOSIT} ke beech hona chahiye.")
        return
    pend = await fetchone("SELECT COUNT(*) c FROM deposits WHERE user_id=? AND status='pending' AND created>?",
                          (m.from_user.id, now() - 1800))
    if pend["c"] >= 3:
        await m.answer("⚠️ Aapke 3 payment pending hain. Pehle unhe complete ya cancel karo.")
        return
    await state.clear()
    oid = await new_order_id()
    cur = await run("INSERT INTO deposits(user_id,method,amount,trx,created) VALUES(?,?,?,?,?)",
                    (m.from_user.id, "upi", amount, oid, now()))
    did = cur.lastrowid
    await send_pay_screen(m, did, oid, amount)
    spawn(poll_deposit(did))
    u = await get_user(m.from_user.id)
    await notify_admins(f"🔔 <b>New Deposit Request</b>\n━━━━━━━━━━━━━━\n👤 {esc(u['name'])} (@{esc(u['username'] or '-')})\n"
                        f"🆔 <code>{u['id']}</code>\n💰 Amount: <b>₹{amount}</b>\n🧾 Order ID: <code>{oid}</code>\n"
                        f"⏳ Status: Payment ka wait (auto verify chal raha hai)")


async def send_pay_screen(m: Message, did, oid, amount):
    upi = SET.get("upi_id", "")
    uri = f"upi://pay?pa={upi}&pn=Payee&am={amount}&cu=INR&tn={oid}"
    cap = (f"{em('deposit')} <b>UPI Payment</b>\n━━━━━━━━━━━━━━\n🧾 Order ID: <code>{oid}</code>\n"
           f"💰 Amount: <b>₹{amount}</b>\n🏦 UPI ID: <code>{esc(upi)}</code>\n\n"
           f"✅ QR ko GPay / PhonePe / Paytm se scan karo (QR ko gallery se bhi scan kar sakte ho).\n"
           f"⚠️ Exact <b>₹{amount}</b> hi bhejo aur payment note me Order ID rehne do.\n"
           f"⏳ Payment hote hi balance <b>apne aap</b> add ho jayega. Na ho to <b>Check Payment</b> dabao.")
    kb = IM([[ib("chkpay", callback_data=f"dc:{did}")],
             [copy_btn("📋 UPI ID", upi), copy_btn("📋 Order ID", oid)],
             [ib("cancelpay", callback_data=f"dx:{did}")]])
    sent = None
    try:
        png = make_qr_png(uri)
        if png:
            sent = await m.answer_photo(BufferedInputFile(png, "upi_qr.png"), caption=cap, reply_markup=kb)
        else:
            from urllib.parse import quote
            sent = await m.answer_photo("https://api.qrserver.com/v1/create-qr-code/?size=400x400&data=" + quote(uri),
                                        caption=cap, reply_markup=kb)
    except Exception:
        logging.exception("qr send failed")
        sent = await m.answer(cap + f"\n\n🔗 UPI Link: <code>{esc(uri)}</code>", reply_markup=kb)
    if sent:
        PAYMSG[did] = (sent.chat.id, sent.message_id)


def spawn(coro):
    t = asyncio.create_task(coro)
    TASKS.add(t)
    t.add_done_callback(TASKS.discard)


async def finish_pay_msg(did, text):
    ref = PAYMSG.pop(did, None)
    if ref:
        try:
            await bot.edit_message_caption(chat_id=ref[0], message_id=ref[1], caption=text, reply_markup=None)
        except Exception:
            pass


async def credit_deposit(did, by="Auto"):
    cur = await DB.execute("UPDATE deposits SET status='approved', decided=? WHERE id=? AND status IN ('pending','expired')",
                           (now(), did))
    if cur.rowcount == 0:
        await DB.commit()
        return False
    d = await fetchone("SELECT * FROM deposits WHERE id=?", (did,))
    await DB.execute("UPDATE users SET balance=balance+? WHERE id=?", (d["amount"], d["user_id"]))
    await DB.commit()
    u = await get_user(d["user_id"])
    await finish_pay_msg(did, f"✅ <b>Payment Received!</b>\n🧾 {d['trx']}\n💰 ₹{d['amount']:g} balance me add ho gaya.")
    try:
        await bot.send_message(d["user_id"], f"✅ <b>Deposit Successful!</b>\n+{money(d['amount'])} aapke balance me add ho gaya.\n"
                                             f"💳 Balance: <b>{money(u['balance'])}</b>\nAb <b>Buy Products</b> se shopping karo 🛒")
    except Exception:
        pass
    await notify_admins(f"✅ <b>Deposit Success ({by})</b>\n━━━━━━━━━━━━━━\n👤 {esc(u['name'])} (@{esc(u['username'] or '-')})\n"
                        f"🆔 <code>{u['id']}</code>\n💰 Amount: <b>₹{d['amount']:g}</b>\n🧾 <code>{d['trx']}</code>")
    return True


async def poll_deposit(did):
    bad, warned = 0, False
    end = time.time() + POLL_MINUTES * 60
    await asyncio.sleep(6)
    while time.time() < end:
        d = await fetchone("SELECT * FROM deposits WHERE id=?", (did,))
        if not d or d["status"] != "pending":
            return
        try:
            st, raw = await vc_call(d["trx"], d["amount"])
        except Exception:
            logging.exception("vc verify error")
            st, raw = "ERROR", ""
        if st == "SUCCESS":
            await credit_deposit(did)
            return
        if st in ("UNKNOWN", "AMOUNT_MISMATCH") and not warned:
            warned = True
            await notify_admins(f"⚠️ <b>Gateway response check karo</b> ({st})\nOrder: <code>{d['trx']}</code>\n"
                                f"<code>{esc(raw[:300])}</code>\nGateway se payment aaya ho to Deposit Requests se manual Approve kar do.")
        if st == "INVALID":
            bad += 1
            if bad >= 3:
                await notify_admins("⚠️ <b>Gateway INVALID bata raha hai.</b> API Key / UPI ID check karo "
                                    "(Payment Settings).")
                return
        await asyncio.sleep(POLL_EVERY)
    cur = await run("UPDATE deposits SET status='expired', decided=? WHERE id=? AND status='pending'", (now(), did))
    if cur.rowcount:
        d = await fetchone("SELECT * FROM deposits WHERE id=?", (did,))
        await finish_pay_msg(did, f"⌛ <b>Payment time khatam</b>\n🧾 {d['trx']}\nAgar paise kat gaye hain to Support me Order ID bhejo.")
        try:
            await bot.send_message(d["user_id"], f"⌛ Order <code>{d['trx']}</code> ka time khatam ho gaya.\n"
                                                 f"Agar aapne payment kar diya hai to <b>Support</b> me Order ID bhejo, hum check karenge.")
        except Exception:
            pass


STATUS_MSG = {
    "PENDING": "⏳ Payment abhi verify ho raha hai. 1-2 minute baad dobara dabao.",
    "NOT_FOUND": "⏳ Payment abhi receive nahi hua. Payment karne ke 1-2 minute baad dobara check karo.",
    "FAILED": "❌ Gateway me payment failed dikh raha hai. Payment dobara try karo ya Support se baat karo.",
    "INVALID": "⚠️ Gateway error. Admin ko batao.",
}


@r.callback_query(F.data.startswith("dc:"))
async def cb_dep_check(c: CallbackQuery):
    did = int(c.data.split(":")[1])
    d = await fetchone("SELECT * FROM deposits WHERE id=?", (did,))
    if not d or (d["user_id"] != c.from_user.id and not is_admin(c.from_user.id)):
        await c.answer("Order nahi mila.", show_alert=True)
        return
    if d["status"] == "approved":
        await c.answer("✅ Ye payment pehle hi add ho chuka hai.", show_alert=True)
        return
    if d["status"] not in ("pending", "expired") or d["created"] < now() - CHECK_WINDOW_HOURS * 3600:
        await c.answer("Ye order ab check nahi ho sakta.", show_alert=True)
        return
    if time.time() - LAST_CHECK.get(c.from_user.id, 0) < 5:
        await c.answer("⏳ 5 second ruko, phir dabao.")
        return
    LAST_CHECK[c.from_user.id] = time.time()
    try:
        st, raw = await vc_call(d["trx"], d["amount"])
    except Exception:
        await c.answer("⚠️ Gateway se connect nahi ho paya. Thodi der baad try karo.", show_alert=True)
        return
    if st == "SUCCESS":
        ok = await credit_deposit(did)
        await c.answer("✅ Payment mil gaya! Balance add ho gaya." if ok else "✅ Pehle hi add ho chuka hai.", show_alert=True)
    else:
        await c.answer(STATUS_MSG.get(st, "⏳ Payment abhi verify nahi hua. Thodi der baad check karo."), show_alert=True)


@r.callback_query(F.data.startswith("dx:"))
async def cb_dep_cancel(c: CallbackQuery):
    did = int(c.data.split(":")[1])
    d = await fetchone("SELECT * FROM deposits WHERE id=?", (did,))
    if not d or d["user_id"] != c.from_user.id:
        await c.answer("Order nahi mila.", show_alert=True)
        return
    cur = await run("UPDATE deposits SET status='cancelled', decided=? WHERE id=? AND status='pending'", (now(), did))
    if cur.rowcount:
        await c.answer("🚫 Cancelled")
        await finish_pay_msg(did, f"🚫 <b>Cancelled</b>\n🧾 {d['trx']}\nAgar paise bhej diye hain to Support me Order ID bhejo.")
    else:
        await c.answer("Ye order ab cancel nahi ho sakta.", show_alert=True)


# ======================================================================
#                 HISTORY (buy / recharge / track order / 2FA)
# ======================================================================
async def h_getcode(m, state):
    kb = IM([[ib("buyhist", callback_data="bh"), ib("rechist", callback_data="rh")],
             [ib("track", callback_data="tr")], [ib("twofa", callback_data="tf")]])
    await m.answer(f"{em('getcode')} <b>History</b>\nNeeche se option chuno:", reply_markup=kb)


@r.callback_query(F.data == "bh")
async def cb_bh(c: CallbackQuery):
    await c.answer()
    uid = c.from_user.id
    t0, t1 = day_bounds()
    q = "SELECT COALESCE(SUM(qty),0) q, COALESCE(SUM(total),0) s FROM orders WHERE user_id=?"
    tot = await fetchone(q, (uid,))
    td = await fetchone(q + " AND created>=?", (uid, t0))
    yd = await fetchone(q + " AND created>=? AND created<?", (uid, t1, t0))
    last = await fetchall("SELECT * FROM orders WHERE user_id=? ORDER BY id DESC LIMIT 10", (uid,))
    lines = "\n".join(f"• <code>{o['invoice']}</code> {esc(o['product_name'])} ×{o['qty']} — ₹{o['total']:g} ({fdate(o['created'])})"
                      for o in last) or "—"
    await c.message.answer(
        f"{em('buyhist')} <b>Buy History</b>\n━━━━━━━━━━━━━━\n"
        f"📅 Aaj: <b>{td['q']} pcs</b> | ₹{td['s']:g}\n📅 Kal: <b>{yd['q']} pcs</b> | ₹{yd['s']:g}\n"
        f"📦 Total: <b>{tot['q']} pcs</b> | ₹{tot['s']:g}\n━━━━━━━━━━━━━━\n<b>Last 10 Orders:</b>\n{lines}")


ICON = {"approved": "✅", "pending": "⏳", "rejected": "❌", "expired": "⌛", "cancelled": "🚫"}


@r.callback_query(F.data == "rh")
async def cb_rh(c: CallbackQuery):
    await c.answer()
    uid = c.from_user.id
    t0, t1 = day_bounds()
    q = "SELECT COALESCE(SUM(amount),0) s FROM deposits WHERE user_id=? AND status='approved'"
    tot = await fetchone(q, (uid,))
    td = await fetchone(q + " AND created>=?", (uid, t0))
    yd = await fetchone(q + " AND created>=? AND created<?", (uid, t1, t0))
    last = await fetchall("SELECT * FROM deposits WHERE user_id=? ORDER BY id DESC LIMIT 10", (uid,))
    lines = "\n".join(f"• {ICON.get(d['status'], '•')} ₹{d['amount']:g} ({fdate(d['created'])})" for d in last) or "—"
    await c.message.answer(
        f"{em('rechist')} <b>Recharge History</b>\n━━━━━━━━━━━━━━\n"
        f"📅 Aaj ka Recharge: <b>₹{td['s']:g}</b>\n📅 Kal ka Recharge: <b>₹{yd['s']:g}</b>\n"
        f"💰 Total Recharge: <b>₹{tot['s']:g}</b>\n━━━━━━━━━━━━━━\n<b>Last 10 Deposits:</b>\n{lines}")


@r.callback_query(F.data == "tr")
async def cb_track(c: CallbackQuery, state: FSMContext):
    await c.answer()
    await state.set_state(S.track)
    await c.message.answer("🔎 <b>Apna Invoice ID likho</b> (jaise <code>INV1791698880451</code>):")


@r.message(S.track, F.text)
async def st_track(m: Message, state: FSMContext):
    await state.clear()
    inv = m.text.strip()
    if is_admin(m.from_user.id):
        o = await fetchone("SELECT * FROM orders WHERE lower(invoice)=lower(?)", (inv,))
    else:
        o = await fetchone("SELECT * FROM orders WHERE lower(invoice)=lower(?) AND user_id=?", (inv, m.from_user.id))
    if not o:
        await m.answer("❌ <b>Ye Invoice ID nahi mila.</b>\nInvoice ID sahi se copy karke daalo.")
        return
    await m.answer(
        f"✅ <b>DELIVERY SUCCESSFUL!</b>\n━━━━━━━━━━━━━━\n🧾 Invoice: <code>{o['invoice']}</code>\n"
        f"📁 Category: <b>{esc(o['cat_name'] or '-')}</b>\n📌 Package: <b>{esc(o['product_name'])}</b>\n"
        f"🔢 Quantity: <b>{o['qty']}</b>\n💰 Amount: <b>₹{o['total']:g}</b>\n📅 Time: {fdate(o['created'])}\n"
        f"━━━━━━━━━━━━━━\nAapka order deliver ho chuka hai.",
        reply_markup=IM([[dib("Login Details dobara dekho", f"trd:{o['id']}", default="📦", style="success")]]))


@r.callback_query(F.data.startswith("trd:"))
async def cb_track_details(c: CallbackQuery):
    await c.answer()
    o = await fetchone("SELECT * FROM orders WHERE id=?", (int(c.data.split(":")[1]),))
    if not o or (o["user_id"] != c.from_user.id and not is_admin(c.from_user.id)):
        return
    await send_items(c.message, [x for x in (o["items"] or "").split("\n") if x.strip()])


@r.callback_query(F.data == "tf")
async def cb_tf(c: CallbackQuery, state: FSMContext):
    await c.answer()
    await state.set_state(S.twofa)
    await c.message.answer("🔐 <b>Apni 2FA Secret Key bhejo:</b>\n(Menu button dabaoge to bahar aa jaoge)")


@r.message(S.twofa, F.text)
async def st_twofa(m: Message):
    try:
        code, left = totp(m.text)
        await m.answer(f"🔐 <b>2FA Code:</b> <code>{code}</code>\n⏳ {left} second valid")
    except Exception:
        await m.answer("❌ <b>Secret Key galat hai.</b>")


# ======================================================================
#                         SUPPORT / TUTORIAL
# ======================================================================
async def h_support(m, state):
    await state.set_state(S.support)
    su = (SET.get("support_user") or "").strip().lstrip("@")
    kb = IM([[ib("contact", url=f"https://t.me/{su}")]]) if su else None
    await m.answer(f"{em('support')} <b>Support Center</b>\n━━━━━━━━━━━━━━\n"
                   f"Apni problem ya sawal detail me likh kar bhejo. Admin bot ke through reply karega."
                   + (f"\n\nDirect contact: @{esc(su)}" if su else ""), reply_markup=kb)


@r.message(S.support)
async def st_support(m: Message):
    u = m.from_user
    head = f"📩 <b>Support Message</b>\n👤 {esc(u.full_name)} (@{esc(u.username or '-')})\n🆔 ID: <code>{u.id}</code>"
    kb = IM([[ib("reply", callback_data=f"sr:{u.id}")]])
    for a in admin_ids():
        try:
            h = await bot.send_message(a, head, reply_markup=kb)
            cp = await bot.copy_message(a, m.chat.id, m.message_id)
            SUPMAP[(a, h.message_id)] = u.id
            SUPMAP[(a, cp.message_id)] = u.id
        except Exception:
            pass
    if len(SUPMAP) > 5000:
        SUPMAP.clear()
    await m.answer("✅ <b>Aapka message admin ko bhej diya gaya.</b> Thoda wait karo, jaldi reply milega.")


async def send_admin_reply(m: Message, uid):
    head = f"{em('a_admins')} <b>Reply from Admin ({esc(m.from_user.full_name)})</b>"
    try:
        if m.text:
            await bot.send_message(uid, f"{head}\n━━━━━━━━━━━━━━\n{esc(m.text)}")
        else:
            await bot.send_message(uid, head)
            await bot.copy_message(uid, m.chat.id, m.message_id)
        await m.answer("✅ Reply bhej diya.")
    except Exception:
        await m.answer("❌ User ko message nahi ja paya (bot block kiya hoga).")


@r.callback_query(A, F.data.startswith("sr:"))
async def cb_sr(c: CallbackQuery, state: FSMContext):
    await c.answer()
    await state.set_state(S.a_reply)
    await state.update_data(uid=int(c.data.split(":")[1]))
    await c.message.answer("↩️ <b>User ke liye reply likho:</b>")


async def h_tutorial(m, state):
    t = SET.get("tutorial", "").strip()
    await m.answer(f"{em('tutorial')} <b>Tutorial</b>\n━━━━━━━━━━━━━━\n{esc(t) if t else 'Tutorial abhi set nahi hai.'}")


# ======================================================================
#                           ADMIN PANEL
# ======================================================================
async def h_admin(m, state):
    await m.answer(f"{em('admin')} <b>Welcome Admin</b>", reply_markup=admin_kb())


async def h_usermenu(m, state):
    await m.answer("✅ <b>Main Menu:</b>", reply_markup=main_kb(m.from_user.id))


async def h_a_prod(m, state):
    kb = IM([[ib("addcat", callback_data="pcm:addcat"), ib("addprod", callback_data="pcat:ap")],
             [ib("chprice", callback_data="pcat:cp")], [ib("delcat", callback_data="pcat:dc")],
             [ib("delprod", callback_data="pcat:dp")], [ib("delstock", callback_data="pcat:ds")]])
    await m.answer(f"{em('a_prod')} <b>Product Management:</b>", reply_markup=kb)


async def pick_cat_kb(action):
    cats = await fetchall("SELECT * FROM categories ORDER BY id")
    rows = []
    for i in range(0, len(cats), 2):
        rows.append([dib(c["name"], f"pc:{action}:{c['id']}", c["emoji_id"]) for c in cats[i:i + 2]])
    if not rows:
        rows = [[dib("Koi category nahi", "noop", default="❌", style="danger")]]
    return IM(rows)


async def h_a_addstock(m, state):
    await m.answer(f"{em('a_addstock')} <b>Kis category me stock add karna hai?</b>", reply_markup=await pick_cat_kb("as"))


@r.callback_query(A, F.data.startswith("pcat:"))
async def cb_pcat(c: CallbackQuery):
    await c.answer()
    await show(c, "📁 <b>Category select karo:</b>", await pick_cat_kb(c.data.split(":")[1]))


@r.callback_query(A, F.data == "pcm:addcat")
async def cb_addcat(c: CallbackQuery, state: FSMContext):
    await c.answer()
    await state.set_state(S.a_cat_name)
    await c.message.answer("📁 <b>Nayi category ka naam likho:</b>")


@r.message(A, S.a_cat_name, F.text)
async def st_cat_name(m: Message, state: FSMContext):
    await state.clear()
    await run("INSERT INTO categories(name) VALUES(?)", (m.text.strip(),))
    await m.answer(f"✅ <b>Category Added:</b> {esc(m.text.strip())}")


@r.callback_query(A, F.data.startswith("pc:"))
async def cb_pc(c: CallbackQuery, state: FSMContext):
    await c.answer()
    _, act, cid = c.data.split(":")
    cid = int(cid)
    cat = await fetchone("SELECT * FROM categories WHERE id=?", (cid,))
    if not cat:
        return
    if act == "ap":
        await state.set_state(S.a_prod_name)
        await state.update_data(cid=cid)
        await c.message.answer(f"📦 <b>{esc(cat['name'])}</b> — product ka naam likho:")
    elif act == "dc":
        await show(c, f"⚠️ <b>{esc(cat['name'])}</b> category, uske saare products aur stock delete ho jayenge. Pakka?",
                   IM([[ib("yes", callback_data=f"cf:dc:{cid}"), ib("no", callback_data="cf:no:0")]]))
    elif act == "emc":
        await state.set_state(S.a_emoji)
        await state.update_data(t="c", k=cid)
        await c.message.answer(f"✨ <b>{esc(cat['name'])}</b> — Premium Emoji ID bhejo (ya premium emoji bhejo). '0' = remove")
    else:
        prods = await fetchall("SELECT * FROM products WHERE cat_id=? ORDER BY id", (cid,))
        rows = []
        for p in prods:
            n = await stock_count(p["id"])
            rows.append([dib(f"{p['name']} | ₹{p['price']:g} | Stock: {n}", f"pp:{act}:{p['id']}", p["emoji_id"], "🌐")])
        if not rows:
            rows = [[dib("Koi product nahi", "noop", default="❌", style="danger")]]
        await show(c, "🌐 <b>Product select karo:</b>", IM(rows))


@r.message(A, S.a_prod_name, F.text)
async def st_prod_name(m: Message, state: FSMContext):
    await state.update_data(name=m.text.strip())
    await state.set_state(S.a_prod_price)
    await m.answer(f"✅ Product: <b>{esc(m.text.strip())}</b>\nAb Normal Price (₹) likho:")


@r.message(A, S.a_prod_price, F.text)
async def st_prod_price(m: Message, state: FSMContext):
    p = parse_amount(m.text, 10_000_000)
    if p is None:
        await m.answer("❌ Number likho.")
        return
    await state.update_data(price=p)
    await state.set_state(S.a_prod_rprice)
    await m.answer(f"✅ Normal Price: <b>₹{p:g}</b>\nAb Reseller Price likho:")


@r.message(A, S.a_prod_rprice, F.text)
async def st_prod_rprice(m: Message, state: FSMContext):
    rp = parse_amount(m.text, 10_000_000)
    if rp is None:
        await m.answer("❌ Number likho.")
        return
    d = await state.get_data()
    await state.clear()
    await run("INSERT INTO products(cat_id,name,price,rprice) VALUES(?,?,?,?)", (d["cid"], d["name"], d["price"], rp))
    await m.answer(f"✅ <b>Product Added!</b>\nName: {esc(d['name'])}\nPrice: ₹{d['price']:g}\nReseller Price: ₹{rp:g}")


@r.callback_query(A, F.data.startswith("pp:"))
async def cb_pp(c: CallbackQuery, state: FSMContext):
    await c.answer()
    _, act, pid = c.data.split(":")
    pid = int(pid)
    p = await fetchone("SELECT * FROM products WHERE id=?", (pid,))
    if not p:
        return
    await state.update_data(pid=pid)
    if act == "as":
        await state.set_state(S.a_stock)
        await c.message.answer(
            f"📦 <b>{esc(p['name'])}</b> ke liye stock bhejo — <b>.xlsx / .xls / .txt / .csv</b> file ya seedha message me.\n\n"
            f"📌 <b>Format (sirf 3 cheez):</b> Link, Username, Password — ek line me ek account:\n"
            f"<code>https://site.com/login | username | password</code>\n\n"
            f"📗 <b>Excel me:</b> Column A = Link, B = Username, C = Password (har row = 1 account).\n"
            f"💡 Ek hi link sabke liye ho to pehli row me likho: <code>Login Link (https://site.com/login)</code> "
            f"aur neeche A = Username, B = Password.")
    elif act == "cp":
        await state.set_state(S.a_price)
        await c.message.answer(f"💲 <b>{esc(p['name'])}</b>\nAbhi: ₹{p['price']:g} | Reseller: ₹{p['rprice']:g}\nNaya Normal Price likho:")
    elif act == "dp":
        await show(c, f"⚠️ <b>{esc(p['name'])}</b> aur uska saara stock delete hoga. Pakka?",
                   IM([[ib("yes", callback_data=f"cf:dp:{pid}"), ib("no", callback_data="cf:no:0")]]))
    elif act == "ds":
        n = await stock_count(pid)
        await show(c, f"⚠️ <b>{esc(p['name'])}</b> ka <b>{n}</b> unsold stock delete hoga. Pakka?",
                   IM([[ib("yes", callback_data=f"cf:ds:{pid}"), ib("no", callback_data="cf:no:0")]]))
    elif act == "emp":
        await state.set_state(S.a_emoji)
        await state.update_data(t="p", k=pid)
        await c.message.answer(f"✨ <b>{esc(p['name'])}</b> — Premium Emoji ID bhejo (ya premium emoji bhejo). '0' = remove")


@r.callback_query(A, F.data.startswith("cf:"))
async def cb_confirm(c: CallbackQuery):
    await c.answer()
    _, act, i = c.data.split(":")
    if act == "no":
        await show(c, "❌ Cancelled")
        return
    i = int(i)
    if act == "dc":
        await DB.execute("DELETE FROM stock WHERE product_id IN (SELECT id FROM products WHERE cat_id=?)", (i,))
        await DB.execute("DELETE FROM products WHERE cat_id=?", (i,))
        await DB.execute("DELETE FROM categories WHERE id=?", (i,))
        msg = "✅ Category Deleted"
    elif act == "dp":
        await DB.execute("DELETE FROM stock WHERE product_id=?", (i,))
        await DB.execute("DELETE FROM products WHERE id=?", (i,))
        msg = "✅ Product Deleted"
    else:
        await DB.execute("DELETE FROM stock WHERE product_id=? AND sold=0", (i,))
        msg = "✅ Stock Deleted"
    await DB.commit()
    await show(c, msg)


@r.message(A, S.a_price, F.text)
async def st_price(m: Message, state: FSMContext):
    p = parse_amount(m.text, 10_000_000)
    if p is None:
        await m.answer("❌ Number likho.")
        return
    await state.update_data(price=p)
    await state.set_state(S.a_rprice)
    await m.answer(f"✅ Normal Price: <b>₹{p:g}</b>\nAb naya Reseller Price likho:")


@r.message(A, S.a_rprice, F.text)
async def st_rprice(m: Message, state: FSMContext):
    rp = parse_amount(m.text, 10_000_000)
    if rp is None:
        await m.answer("❌ Number likho.")
        return
    d = await state.get_data()
    await state.clear()
    await run("UPDATE products SET price=?, rprice=? WHERE id=?", (d["price"], rp, d["pid"]))
    await m.answer(f"✅ <b>Price Updated!</b>\nNormal: ₹{d['price']:g}\nReseller: ₹{rp:g}")


@r.message(A, S.a_stock)
async def st_stock(m: Message, state: FSMContext):
    d = await state.get_data()
    items = []
    try:
        if m.document:
            buf = io.BytesIO()
            await bot.download(m.document, destination=buf)
            items = await asyncio.to_thread(parse_stock_file, m.document.file_name, buf.getvalue())
        elif m.text:
            items = parse_txt(m.text.encode())
    except Exception as e:
        await m.answer(f"❌ File padh nahi paya: {esc(str(e))}\n.xlsx ya .txt me bhejo.")
        return
    if not items:
        await m.answer("❌ Koi item nahi mila. .xlsx / .xls / .txt file ya text bhejo.")
        return
    await state.clear()
    await DB.executemany("INSERT INTO stock(product_id,data) VALUES(?,?)", [(d["pid"], x) for x in items])
    await DB.commit()
    p = await fetchone("SELECT name FROM products WHERE id=?", (d["pid"],))
    await m.answer(f"✅ <b>Stock Added!</b>\n📦 {esc(p['name'])}\n➕ Added: <b>{len(items)}</b>\n🔥 Total Stock: <b>{await stock_count(d['pid'])}</b>")


async def h_a_live(m, state):
    cats = await fetchall("SELECT * FROM categories ORDER BY id")
    if not cats:
        await m.answer("❌ Koi category nahi hai.")
        return
    out = f"{em('a_live')} <b>Live Stock</b>\n"
    for cat in cats:
        out += f"\n📁 <b>{esc(cat['name'])}</b>\n"
        prods = await fetchall("SELECT * FROM products WHERE cat_id=?", (cat["id"],))
        for p in prods:
            n = await stock_count(p["id"])
            out += f"  🌐 <b>{esc(p['name'])}</b> | ₹{p['price']:g} (Reseller: ₹{p['rprice']:g}) | Stock: <b>{n}</b>\n"
        if not prods:
            out += "  —\n"
    for i in range(0, len(out), 3800):
        await m.answer(out[i:i + 3800])


# ---------------- user management ----------------
async def h_a_users(m, state):
    await state.set_state(S.a_user)
    await m.answer("👥 <b>User ka Telegram ID ya Username do:</b>")


async def user_card_view(uid):
    u = await get_user(uid)
    o = await fetchone("SELECT COALESCE(SUM(qty),0) q FROM orders WHERE user_id=?", (uid,))
    txt = (f"👥 <b>USER PROFILE</b>\n━━━━━━━━━━━━━━\n👑 Name: <b>{esc(u['name'])}</b>\n⚙️ ID: <code>{u['id']}</code>\n"
           f"🌐 Username: @{esc(u['username'] or '-')}\n💰 Balance: <b>{money(u['balance'])}</b>\n"
           f"🧾 Orders: <b>{o['q']}</b>\n🤝 Reseller: <b>{'Yes' if u['reseller'] else 'No'}</b>\n"
           f"Status: <b>{'🚫 Banned' if u['banned'] else '✅ Active'}</b>\n━━━━━━━━━━━━━━")
    kb = IM([[ib("unban", callback_data=f"um:unban:{uid}") if u["banned"] else ib("ban", callback_data=f"um:ban:{uid}")],
             [ib("addbal", callback_data=f"um:add:{uid}"), ib("cutbal", callback_data=f"um:cut:{uid}")],
             [ib("sendmsg", callback_data=f"um:msg:{uid}")]])
    return txt, kb


@r.message(A, S.a_user, F.text)
async def st_a_user(m: Message, state: FSMContext):
    u = await find_user(m.text)
    if not u:
        await m.answer("❌ User nahi mila (user ne ek baar bot start kiya hona chahiye).")
        return
    await state.clear()
    t, kb = await user_card_view(u["id"])
    await m.answer(t, reply_markup=kb)


@r.callback_query(A, F.data.startswith("um:"))
async def cb_um(c: CallbackQuery, state: FSMContext):
    await c.answer()
    _, act, uid = c.data.split(":")
    uid = int(uid)
    if act in ("ban", "unban"):
        await run("UPDATE users SET banned=? WHERE id=?", (1 if act == "ban" else 0, uid))
        t, kb = await user_card_view(uid)
        await show(c, t, kb)
    elif act in ("add", "cut"):
        await state.set_state(S.a_bal)
        await state.update_data(uid=uid, mode=act)
        await c.message.answer(f"💰 Kitne ₹ {'Add' if act == 'add' else 'Cut'} karne hain?")
    elif act == "msg":
        await state.set_state(S.a_msg)
        await state.update_data(uid=uid)
        await c.message.answer("✉️ User ko jo message bhejna hai wo bhejo:")


@r.message(A, S.a_bal, F.text)
async def st_a_bal(m: Message, state: FSMContext):
    amt = parse_amount(m.text, 10_000_000)
    if amt is None:
        await m.answer("❌ Sahi number likho.")
        return
    d = await state.get_data()
    await state.clear()
    sign = 1 if d["mode"] == "add" else -1
    await run("UPDATE users SET balance=balance+? WHERE id=?", (sign * amt, d["uid"]))
    await m.answer(f"✅ <b>{money(amt)}</b> {'add' if sign > 0 else 'cut'} ho gaya.")
    try:
        await bot.send_message(d["uid"], f"{'✅ +' if sign > 0 else '➖ -'}<b>{money(amt)}</b> aapke balance me {'add' if sign > 0 else 'cut'} hua.")
    except Exception:
        pass


@r.message(A, S.a_msg)
async def st_a_msg(m: Message, state: FSMContext):
    d = await state.get_data()
    await state.clear()
    try:
        await bot.copy_message(d["uid"], m.chat.id, m.message_id)
        await m.answer("✅ Message bhej diya.")
    except Exception:
        await m.answer("❌ Message nahi ja paya.")


@r.message(A, S.a_reply)
async def st_a_reply(m: Message, state: FSMContext):
    d = await state.get_data()
    await state.clear()
    await send_admin_reply(m, d["uid"])


@r.message(A, SupReplyF())
async def sup_reply_by_reply(m: Message, sup_uid: int):
    await send_admin_reply(m, sup_uid)


# ---------------- deposits (admin) ----------------
async def h_a_depreq(m, state):
    rows = await fetchall("SELECT * FROM deposits WHERE status IN ('pending','expired') AND created>? ORDER BY id DESC LIMIT 15",
                          (now() - CHECK_WINDOW_HOURS * 3600,))
    if not rows:
        await m.answer("❌ <b>Koi pending deposit nahi hai.</b>")
        return
    for d in rows:
        u = await get_user(d["user_id"])
        await m.answer(
            f"🚨 <b>Deposit #{d['id']}</b> ({ICON.get(d['status'])} {d['status']})\n━━━━━━━━━━━━━━\n"
            f"👤 {esc(u['name'])} (@{esc(u['username'] or '-')})\n🆔 <code>{u['id']}</code>\n"
            f"💰 <b>₹{d['amount']:g}</b>\n🧾 <code>{d['trx']}</code>\n📅 {fdate(d['created'])}\n\n"
            f"⚠️ Sirf tab Approve karo jab aapke UPI/Paytm me payment dikh raha ho.",
            reply_markup=IM([[ib("approve", callback_data=f"da:{d['id']}"), ib("reject", callback_data=f"dr:{d['id']}")]]))


@r.callback_query(A, F.data.regexp(r"^d[ar]:\d+$"))
async def cb_dep_decide(c: CallbackQuery):
    act, did = c.data.split(":")
    did = int(did)
    d = await fetchone("SELECT * FROM deposits WHERE id=?", (did,))
    if not d:
        await c.answer("Not found", show_alert=True)
        return
    if act == "da":
        ok = await credit_deposit(did, f"Manual by {c.from_user.full_name}")
        await c.answer("✅ Approved" if ok else "⚠️ Pehle hi process ho chuka hai.", show_alert=not ok)
        if ok:
            await show(c, f"✅ <b>Deposit #{did} Approved</b> — ₹{d['amount']:g}")
    else:
        cur = await run("UPDATE deposits SET status='rejected', decided=? WHERE id=? AND status IN ('pending','expired')", (now(), did))
        if cur.rowcount == 0:
            await c.answer("⚠️ Pehle hi process ho chuka hai.", show_alert=True)
            return
        await c.answer("Rejected")
        await show(c, f"❌ <b>Deposit #{did} Rejected</b>")
        await finish_pay_msg(did, f"❌ <b>Rejected</b>\n🧾 {d['trx']}")
        try:
            await bot.send_message(d["user_id"], f"❌ <b>Deposit Rejected:</b> ₹{d['amount']:g}\nOrder: <code>{d['trx']}</code>\n"
                                                 f"Problem ho to Support me contact karo.")
        except Exception:
            pass


async def h_a_dephist(m, state):
    t0, _ = day_bounds()
    q = "SELECT COALESCE(SUM(amount),0) s, COUNT(*) c FROM deposits WHERE status='approved'"
    tot = await fetchone(q)
    td = await fetchone(q + " AND created>=?", (t0,))
    pend = await fetchone("SELECT COUNT(*) c FROM deposits WHERE status='pending'")
    last = await fetchall("SELECT * FROM deposits ORDER BY id DESC LIMIT 20")
    lines = "\n".join(f"{ICON.get(d['status'], '•')} #{d['id']} <code>{d['user_id']}</code> ₹{d['amount']:g} ({fdate(d['created'])})"
                      for d in last) or "—"
    await m.answer(f"{em('a_dephist')} <b>Deposit History</b>\n━━━━━━━━━━━━━━\n📅 Aaj: <b>₹{td['s']:g}</b> ({td['c']})\n"
                   f"💰 Total Approved: <b>₹{tot['s']:g}</b> ({tot['c']})\n⏳ Pending: <b>{pend['c']}</b>\n━━━━━━━━━━━━━━\n{lines}")


# ---------------- broadcast ----------------
async def h_a_bcast(m, state):
    await state.set_state(S.a_bcast)
    await m.answer("📢 <b>Broadcast message bhejo</b> (text/photo/video kuch bhi).\nCancel: /cancel")


@r.message(A, S.a_bcast)
async def st_bcast(m: Message, state: FSMContext):
    await state.clear()
    users = await fetchall("SELECT id FROM users WHERE banned=0")
    st = await m.answer(f"⏳ {len(users)} users ko bhej raha hu...")
    ok = fail = 0
    for u in users:
        try:
            await bot.copy_message(u["id"], m.chat.id, m.message_id)
            ok += 1
        except Exception:
            fail += 1
        await asyncio.sleep(0.04)
    await st.edit_text(f"📢 <b>Broadcast Done</b>\n✅ Success: <b>{ok}</b>\n❌ Failed: <b>{fail}</b>")


# ---------------- vouchers ----------------
async def h_a_voucher(m, state):
    kb = IM([[ib("vcreate", callback_data="v:new")], [ib("vlist", callback_data="v:list")], [ib("vdelete", callback_data="v:del")]])
    await m.answer(f"{em('a_voucher')} <b>Voucher Management:</b>", reply_markup=kb)


@r.callback_query(A, F.data.startswith("v:"))
async def cb_v(c: CallbackQuery, state: FSMContext):
    await c.answer()
    act = c.data.split(":")[1]
    if act == "new":
        await state.set_state(S.a_vcode)
        await c.message.answer("🎟 <b>Voucher Code likho</b> (jaise: WELCOME20):")
    elif act == "list":
        vs = await fetchall("SELECT * FROM vouchers")
        t = "\n".join(f"🎟 <code>{esc(v['code'])}</code> | ₹{v['amount']:g} | Used {v['used']}/{v['max_uses']}" for v in vs) or "Koi voucher nahi."
        await c.message.answer(f"📋 <b>Vouchers</b>\n{t}")
    elif act == "del":
        vs = await fetchall("SELECT * FROM vouchers")
        rows = [[dib(f"{v['code']} (₹{v['amount']:g})", f"vdel:{v['code']}", default="🗑", style="danger")] for v in vs]
        if rows:
            await c.message.answer("🗑 <b>Kaunsa delete karna hai?</b>", reply_markup=IM(rows))
        else:
            await c.message.answer("Koi voucher nahi.")


@r.callback_query(A, F.data.startswith("vdel:"))
async def cb_vdel(c: CallbackQuery):
    await c.answer("Deleted")
    await run("DELETE FROM vouchers WHERE code=?", (c.data.split(":", 1)[1],))
    await show(c, "✅ Voucher Deleted")


@r.message(A, S.a_vcode, F.text)
async def st_vcode(m: Message, state: FSMContext):
    code = m.text.strip().upper()[:20]
    if not re.fullmatch(r"[A-Z0-9_\-]+", code):
        await m.answer("❌ Sirf English letters/numbers use karo.")
        return
    await state.update_data(code=code)
    await state.set_state(S.a_vamt)
    await m.answer("💰 Kitne ₹ ka voucher?")


@r.message(A, S.a_vamt, F.text)
async def st_vamt(m: Message, state: FSMContext):
    amt = parse_amount(m.text, 10_000_000)
    if amt is None:
        await m.answer("❌ Number likho.")
        return
    await state.update_data(amt=amt)
    await state.set_state(S.a_vuses)
    await m.answer("👥 Maximum kitne log use kar sakte hain?")


@r.message(A, S.a_vuses, F.text)
async def st_vuses(m: Message, state: FSMContext):
    if not m.text.strip().isdigit():
        await m.answer("❌ Number likho.")
        return
    d = await state.get_data()
    await state.clear()
    await run("INSERT OR REPLACE INTO vouchers(code,amount,max_uses,used) VALUES(?,?,?,0)", (d["code"], d["amt"], int(m.text.strip())))
    await m.answer(f"✅ <b>Voucher Created</b>\nCode: <code>{d['code']}</code>\nAmount: ₹{d['amt']:g}\nLimit: {m.text.strip()}")


# ---------------- force join ----------------
async def h_a_force(m, state):
    chs = await fetchall("SELECT * FROM force_channels")
    t = "\n".join(f"• {esc(c['title'] or c['chat'])} (<code>{esc(c['chat'])}</code>)" for c in chs) or "Koi channel set nahi hai."
    kb = IM([[ib("fadd", callback_data="f:add")], [ib("fremove", callback_data="f:rm")]])
    await m.answer(f"{em('a_force')} <b>Force Join Setup</b>\n━━━━━━━━━━━━━━\n{t}\n\n⚠️ Bot ko channel/group me Admin banana zaroori hai.", reply_markup=kb)


@r.callback_query(A, F.data.startswith("f:"))
async def cb_f(c: CallbackQuery, state: FSMContext):
    await c.answer()
    if c.data == "f:add":
        await state.set_state(S.a_force)
        await c.message.answer("🔒 Channel ka @username ya ID (-100...) bhejo:")
    else:
        chs = await fetchall("SELECT * FROM force_channels")
        rows = [[dib(x["title"] or x["chat"], f"frm:{x['id']}", default="🗑", style="danger")] for x in chs]
        if rows:
            await c.message.answer("🗑 <b>Kaunsa remove karna hai?</b>", reply_markup=IM(rows))
        else:
            await c.message.answer("Kuch nahi hai.")


@r.callback_query(A, F.data.startswith("frm:"))
async def cb_frm(c: CallbackQuery):
    await c.answer("Removed")
    await run("DELETE FROM force_channels WHERE id=?", (int(c.data.split(":")[1]),))
    await show(c, "✅ Channel Removed")


@r.message(A, S.a_force, F.text)
async def st_force(m: Message, state: FSMContext):
    x = m.text.strip()
    chat = int(x) if x.lstrip("-").isdigit() else x
    try:
        ch = await bot.get_chat(chat)
    except Exception:
        await m.answer("❌ Channel nahi mila. Bot ko admin banaya ya nahi check karo.")
        return
    await state.clear()
    link = f"https://t.me/{ch.username}" if ch.username else (ch.invite_link or "")
    await run("INSERT INTO force_channels(chat,title,link) VALUES(?,?,?)", (str(ch.id), ch.title or x, link))
    await m.answer(f"✅ <b>Added:</b> {esc(ch.title or x)}")


# ---------------- support / tutorial / payment settings ----------------
async def ask_set(m, state, key, prompt):
    await state.set_state(S.a_setval)
    await state.update_data(key=key)
    await m.answer(prompt)


async def h_a_support(m, state):
    await ask_set(m, state, "support_user", f"🎧 <b>Support Setup</b>\nAbhi: @{esc(SET.get('support_user') or '-')}\nNaya Support Username likho:")


async def h_a_tutorial(m, state):
    await ask_set(m, state, "tutorial", "📚 <b>Tutorial Setup</b>\nTutorial ka text/link likho (poora message):")


def mask_key(k):
    return (k[:4] + "****" + k[-3:]) if len(k) > 8 else "****"


async def h_a_pay(m, state):
    mn = SET.get("min_dep")
    kb = IM([[ib("edit_upi", callback_data="ps:upi_id")], [ib("edit_key", callback_data="ps:api_key")],
             [ib("edit_min", f"Edit Minimum Deposit (Abhi: ₹{mn})", callback_data="ps:min_dep")]])
    await m.answer(f"{em('a_pay')} <b>Payment Settings (UPI Auto)</b>\n━━━━━━━━━━━━━━\n🏦 UPI ID: <code>{esc(SET.get('upi_id', ''))}</code>\n"
                   f"🔑 API Key: <code>{esc(mask_key(SET.get('api_key', '')))}</code>\n⬇️ Min Deposit: <b>₹{mn}</b>\n\n"
                   f"💡 Pehle <b>Min Deposit = 1</b> karke ₹1 se test karo, phir 10+ rakh do.", reply_markup=kb)


@r.callback_query(A, F.data.startswith("ps:"))
async def cb_ps(c: CallbackQuery, state: FSMContext):
    await c.answer()
    key = c.data.split(":")[1]
    names = {"upi_id": "UPI ID (jaise name@bank)", "api_key": "Gateway API Key", "min_dep": "Minimum Deposit (₹)"}
    await ask_set(c.message, state, key, f"✏️ Naya <b>{names[key]}</b> likho:")


@r.message(A, S.a_setval, F.text)
async def st_setval(m: Message, state: FSMContext):
    d = await state.get_data()
    key, val = d["key"], m.text.strip()
    if key == "min_dep":
        if not val.isdigit() or int(val) < 1:
            await m.answer("❌ Poora number likho (minimum 1).")
            return
    elif key == "upi_id" and not re.fullmatch(r"[\w.\-]+@[\w.\-]+", val):
        await m.answer("❌ UPI ID sahi format me likho (jaise name@bank).")
        return
    elif key == "support_user":
        val = val.lstrip("@")
    await state.clear()
    await set_setting(key, val)
    if key == "api_key":
        try:
            await m.delete()
        except Exception:
            pass
    await m.answer("✅ <b>Updated!</b>")


# ---------------- channel setup (purchase notification) ----------------
async def h_a_channel(m, state):
    cur = SET.get("notify_chat") or ""
    kb = IM([[dib("Channel Set Karo", "chn:set", default="➕", style="success"), dib("Band Karo", "chn:off", default="🗑", style="danger")],
             [dib("Test Message", "chn:test", default="🧪")]])
    await m.answer(f"{em('a_channel')} <b>Channel Setup</b>\n━━━━━━━━━━━━━━\nHar purchase ke baad is channel me "
                   f"<b>AUTO DELIVERY SUCCESS</b> notification jayega.\nAbhi: <code>{esc(cur) if cur else 'OFF'}</code>\n\n"
                   f"⚠️ Bot ko channel me <b>Admin</b> banao (Post messages permission).", reply_markup=kb)


@r.callback_query(A, F.data.startswith("chn:"))
async def cb_chn(c: CallbackQuery, state: FSMContext):
    await c.answer()
    act = c.data.split(":")[1]
    if act == "set":
        await state.set_state(S.a_channel)
        await c.message.answer("📣 Channel ka @username ya ID (-100...) bhejo:")
    elif act == "off":
        await set_setting("notify_chat", "")
        await c.message.answer("✅ Channel notification band.")
    else:
        if not SET.get("notify_chat"):
            await c.message.answer("❌ Pehle channel set karo.")
            return
        await notify_channel(channel_text("INV1791698880451", "Test User", "testuser12", 123456789, "Category", "Package", 1, 12, now()))
        await c.message.answer("🧪 Test message bhej diya (channel check karo).")


@r.message(A, S.a_channel, F.text)
async def st_channel(m: Message, state: FSMContext):
    x = m.text.strip()
    chat = int(x) if x.lstrip("-").isdigit() else x
    try:
        ch = await bot.get_chat(chat)
        await bot.send_message(ch.id, "✅ <b>Notification channel connected!</b>\nAb har purchase ka update yaha aayega.")
    except Exception:
        await m.answer("❌ Channel me message nahi ja paya. Bot ko Admin banao (Post messages) aur dobara try karo.")
        return
    await state.clear()
    await set_setting("notify_chat", str(ch.id))
    await m.answer(f"✅ <b>Channel set:</b> {esc(ch.title or x)}")


# ---------------- admin / reseller management ----------------
async def h_a_admins(m, state):
    if m.from_user.id not in OWNER_IDS:
        await m.answer("❌ Ye option sirf Owner ke liye hai.")
        return
    t = "\n".join(f"• <code>{a}</code>" for a in sorted(ADMINS)) or "Koi extra admin nahi hai."
    kb = IM([[ib("addadmin", callback_data="adm:add")], [ib("deladmin", callback_data="adm:del")]])
    await m.answer(f"{em('a_admins')} <b>Admin Management System:</b>\n{t}", reply_markup=kb)


@r.callback_query(A, F.data.startswith("adm:"))
async def cb_adm(c: CallbackQuery, state: FSMContext):
    await c.answer()
    if c.from_user.id not in OWNER_IDS:
        return
    if c.data == "adm:add":
        await state.set_state(S.a_admin)
        await c.message.answer("👑 Naye Admin ka Telegram ID likho:")
    else:
        rows = [[dib(str(a), f"admd:{a}", default="❌", style="danger")] for a in sorted(ADMINS)]
        if rows:
            await c.message.answer("Kisko remove karna hai?", reply_markup=IM(rows))
        else:
            await c.message.answer("Koi extra admin nahi hai.")


@r.callback_query(A, F.data.startswith("admd:"))
async def cb_admd(c: CallbackQuery):
    if c.from_user.id not in OWNER_IDS:
        await c.answer()
        return
    await c.answer("Removed")
    await run("DELETE FROM admins WHERE id=?", (int(c.data.split(":")[1]),))
    await load_cache()
    await show(c, "✅ Admin Removed")


@r.message(A, S.a_admin, F.text)
async def st_admin(m: Message, state: FSMContext):
    if m.from_user.id not in OWNER_IDS or not m.text.strip().isdigit():
        await m.answer("❌ Sahi numeric ID do.")
        return
    await state.clear()
    await run("INSERT OR IGNORE INTO admins(id) VALUES(?)", (int(m.text.strip()),))
    await load_cache()
    await m.answer("✅ <b>Admin Added!</b>")


async def h_a_resell(m, state):
    rs = await fetchall("SELECT id,name FROM users WHERE reseller=1")
    t = "\n".join(f"• {esc(x['name'])} <code>{x['id']}</code>" for x in rs) or "Koi reseller nahi hai."
    kb = IM([[ib("addresell", callback_data="rs:add"), ib("delresell", callback_data="rs:del")]])
    await m.answer(f"{em('a_resell')} <b>Reseller Management</b>\n{t}", reply_markup=kb)


@r.callback_query(A, F.data.startswith("rs:"))
async def cb_rs(c: CallbackQuery, state: FSMContext):
    await c.answer()
    await state.set_state(S.a_resell)
    await state.update_data(mode=c.data.split(":")[1])
    await c.message.answer("🤝 User ka Telegram ID ya Username do:")


@r.message(A, S.a_resell, F.text)
async def st_resell(m: Message, state: FSMContext):
    u = await find_user(m.text)
    if not u:
        await m.answer("❌ User nahi mila.")
        return
    d = await state.get_data()
    await state.clear()
    on = 1 if d["mode"] == "add" else 0
    await run("UPDATE users SET reseller=? WHERE id=?", (on, u["id"]))
    await m.answer(f"✅ {esc(u['name'])} ab {'Reseller' if on else 'Normal User'} hai.")
    try:
        await bot.send_message(u["id"], "🤝 <b>Aap ab Reseller ho!</b> Ab reseller price par kharid sakte ho." if on
                               else "ℹ️ Aapka Reseller status hata diya gaya.")
    except Exception:
        pass


# ---------------- premium emoji: buttons ----------------
EMO_KEYS = list(B.keys())


def emoji_page_kb(page):
    per = 14
    chunk = EMO_KEYS[page * per:(page + 1) * per]
    rows = []
    for i in range(0, len(chunk), 2):
        rows.append([InlineKeyboardButton(text=("✅ " if k in EMO else "") + B[k][1], callback_data=f"es:{k}", style="primary")
                     for k in chunk[i:i + 2]])
    nav = []
    if page > 0:
        nav.append(InlineKeyboardButton(text="◀️ Prev", callback_data=f"ek:{page - 1}"))
    if (page + 1) * per < len(EMO_KEYS):
        nav.append(InlineKeyboardButton(text="Next ▶️", callback_data=f"ek:{page + 1}"))
    if nav:
        rows.append(nav)
    rows.append([InlineKeyboardButton(text="📁 Category Emoji", callback_data="pcat:emc", style="success"),
                 InlineKeyboardButton(text="🌐 Product Emoji", callback_data="pcat:emp", style="success")])
    rows.append([InlineKeyboardButton(text="📝 Text Emoji (message ke andar)", callback_data="tx:menu", style="danger")])
    return IM(rows)


async def h_a_emoji(m, state):
    await m.answer(f"{em('a_emoji')} <b>Premium Emoji Set</b>\nJis button me Premium Emoji lagana hai wo chuno (✅ = set hai):",
                   reply_markup=emoji_page_kb(0))


@r.callback_query(A, F.data.startswith("ek:"))
async def cb_ek(c: CallbackQuery):
    await c.answer()
    await show(c, f"{em('a_emoji')} <b>Premium Emoji Set</b>\nJis button me Premium Emoji lagana hai wo chuno:",
               emoji_page_kb(int(c.data.split(":")[1])))


@r.callback_query(A, F.data.startswith("es:"))
async def cb_es(c: CallbackQuery, state: FSMContext):
    await c.answer()
    key = c.data.split(":", 1)[1]
    if key not in B:
        return
    await state.set_state(S.a_emoji)
    await state.update_data(t="b", k=key)
    await c.message.answer(f"✨ <b>{B[key][1]}</b> button ke liye Premium Emoji ID bhejo\n(ya seedha ek premium emoji bhejo). '0' = remove.")


@r.message(A, S.a_emoji, F.text)
async def st_emoji(m: Message, state: FSMContext):
    eid = parse_emoji_id(m)
    if eid is None:
        await m.answer("❌ Sirf Emoji ID (number) ya premium emoji bhejo.")
        return
    d = await state.get_data()
    await state.clear()
    remove = eid == "0"
    if d["t"] == "b":
        if remove:
            await run("DELETE FROM emojis WHERE key=?", (d["k"],))
        else:
            await run("INSERT OR REPLACE INTO emojis(key,emoji_id) VALUES(?,?)", (d["k"], eid))
        await load_cache()
    else:
        table = "categories" if d["t"] == "c" else "products"
        await run(f"UPDATE {table} SET emoji_id=? WHERE id=?", (None if remove else eid, d["k"]))
    await m.answer(f"✅ <b>{'Removed' if remove else 'Set!'}</b> {'' if remove else cem(eid, '⭐')}", reply_markup=admin_kb())


# ---------------- premium emoji: text ----------------
def tx_menu_kb():
    return IM([[dib("Add / Update", "tx:add", default="➕", style="success"), dib("Remove", "tx:del", default="🗑", style="danger")],
               [dib("Mapped List", "tx:list", default="📋")]])


@r.callback_query(A, F.data.startswith("tx:"))
async def cb_tx(c: CallbackQuery, state: FSMContext):
    await c.answer()
    act = c.data.split(":")[1]
    if act == "menu":
        await show(c, "📝 <b>Text Emoji</b>\nBot ke text me jo normal emoji (👑 ⚙️ 🌐 💰 ✅ ...) aate hain unki jagah premium emoji "
                      f"dikhane ke liye yaha set karo.\nAbhi set: <b>{len(TEMO)}</b>", tx_menu_kb())
    elif act == "add":
        await state.set_state(S.a_temoji)
        await c.message.answer(
            "📝 <b>Premium emoji bhejo</b> (ek ya kai, ek hi message me), jaise: 👑 ⚙️ 🌐 💰 🧾\n"
            "Bot khud match kar lega: bot ke text me jo normal emoji hai uski jagah ye premium emoji dikhega.\n\n"
            "Agar premium emoji ka normal roop alag ho to aise likho (ek line me ek):\n"
            "<code>👑 5123456789012345678</code>  (normal emoji + premium ID)")
    elif act == "del":
        await state.set_state(S.a_temoji_del)
        await c.message.answer("🗑 Jis normal emoji ka premium hataana hai wo emoji bhejo (kai bhi bhej sakte ho):")
    elif act == "list":
        if not TEMO:
            await c.message.answer("Koi Text Emoji set nahi hai.")
            return
        await c.message.answer("📋 <b>Mapped Text Emoji:</b>\n" + "  ".join(cem(i, k) for k, i in TEMO.items()))


@r.message(A, S.a_temoji, F.text)
async def st_temoji(m: Message, state: FSMContext):
    pairs = {}
    for e in (m.entities or []):
        if e.type == "custom_emoji":
            ch = _nk(e.extract_from(m.text)).strip()
            if ch:
                pairs[ch] = e.custom_emoji_id
    for line in m.text.splitlines():
        mo = re.match(r"^\s*(\S+)\s+(\d{8,})\s*$", line)
        if mo:
            pairs[_nk(mo.group(1))] = mo.group(2)
    if not pairs:
        await m.answer("❌ Premium emoji nahi mila. Premium emoji bhejo ya <code>👑 5123456789012345678</code> format use karo.")
        return
    await state.clear()
    for ch, i in pairs.items():
        await DB.execute("INSERT OR REPLACE INTO text_emojis(ch,emoji_id) VALUES(?,?)", (ch, i))
    await DB.commit()
    await load_cache()
    await m.answer(f"✅ <b>{len(pairs)} Text Emoji Set!</b>\n" + "  ".join(cem(i, ch) for ch, i in pairs.items()))


@r.message(A, S.a_temoji_del, F.text)
async def st_temoji_del(m: Message, state: FSMContext):
    t = _nk(m.text)
    gone = [k for k in TEMO if k in t]
    await state.clear()
    for k in gone:
        await DB.execute("DELETE FROM text_emojis WHERE ch=?", (k,))
    await DB.commit()
    await load_cache()
    await m.answer(f"✅ {len(gone)} Text Emoji hata diye." if gone else "❌ Wo emoji set nahi tha.")


# ---------------- statistics ----------------
async def h_a_stats(m, state):
    u = await fetchone("SELECT COUNT(*) c, COALESCE(SUM(balance),0) b FROM users")
    o = await fetchone("SELECT COALESCE(SUM(qty),0) q, COALESCE(SUM(total),0) s FROM orders")
    d = await fetchone("SELECT COALESCE(SUM(amount),0) s FROM deposits WHERE status='approved'")
    st = await fetchone("SELECT COUNT(*) c FROM stock WHERE sold=0")
    t0, _ = day_bounds()
    to = await fetchone("SELECT COALESCE(SUM(qty),0) q, COALESCE(SUM(total),0) s FROM orders WHERE created>=?", (t0,))
    nu = await fetchone("SELECT COUNT(*) c FROM users WHERE joined>=?", (t0,))
    await m.answer(f"{em('a_stats')} <b>Statistics</b>\n━━━━━━━━━━━━━━\n👥 Users: <b>{u['c']}</b> (aaj naye: {nu['c']})\n"
                   f"💰 User Balances: <b>₹{u['b']:.2f}</b>\n🛍 Total Sold: <b>{o['q']} pcs</b> | ₹{o['s']:.2f}\n"
                   f"📅 Aaj ki Sale: <b>{to['q']} pcs</b> | ₹{to['s']:.2f}\n💵 Total Deposit: <b>₹{d['s']:.2f}</b>\n📦 Total Stock: <b>{st['c']}</b>")


# ---------------- backup + gateway test commands ----------------
async def make_backup():
    path = os.path.join(BACKUP_DIR, f"shop_{datetime.now(IST).strftime('%Y%m%d_%H%M%S')}.db")
    await DB.commit()
    await DB.execute("VACUUM INTO ?", (path,))
    for old in sorted(f for f in os.listdir(BACKUP_DIR) if f.endswith(".db"))[:-7]:
        try:
            os.remove(os.path.join(BACKUP_DIR, old))
        except OSError:
            pass
    return path


async def send_backup(chat_id, caption="🗄 <b>Database Backup</b>"):
    path = await make_backup()
    with open(path, "rb") as f:
        data = f.read()
    await bot.send_document(chat_id, BufferedInputFile(data, os.path.basename(path)), caption=caption)


async def backup_loop():
    while True:
        await asyncio.sleep(BACKUP_EVERY_HOURS * 3600)
        try:
            for o in OWNER_IDS:
                await send_backup(o, "🗄 <b>Auto Database Backup</b>")
        except Exception:
            logging.exception("auto backup failed")


@r.message(Command("backup"))
async def cmd_backup(m: Message):
    if m.from_user.id in OWNER_IDS:
        await send_backup(m.chat.id)


@r.message(Command("gwcheck"))
async def cmd_gwcheck(m: Message):
    """/gwcheck ORDERID [AMOUNT]  -> gateway ka raw response dikhata hai (debug)"""
    if not is_admin(m.from_user.id):
        return
    parts = (m.text or "").split()
    if len(parts) < 2:
        await m.answer("Use: <code>/gwcheck ORDER_ID [AMOUNT]</code>")
        return
    oid = parts[1]
    amt = parts[2] if len(parts) > 2 else None
    if amt is None:
        d = await fetchone("SELECT amount FROM deposits WHERE trx=?", (oid,))
        amt = d["amount"] if d else 1
    try:
        st, raw = await vc_call(oid, float(amt))
    except Exception as e:
        await m.answer(f"❌ Gateway error: <code>{esc(str(e))}</code>")
        return
    await m.answer(f"🔎 <b>Status:</b> {st}\n<b>Raw response:</b>\n<code>{esc(raw[:700])}</code>")


MENU_H.update({
    "buy": h_buy, "profile": h_profile, "deposit": h_deposit, "getcode": h_getcode, "support": h_support,
    "tutorial": h_tutorial, "admin": h_admin, "usermenu": h_usermenu,
    "a_prod": h_a_prod, "a_addstock": h_a_addstock, "a_live": h_a_live, "a_users": h_a_users,
    "a_depreq": h_a_depreq, "a_dephist": h_a_dephist, "a_bcast": h_a_bcast, "a_voucher": h_a_voucher,
    "a_force": h_a_force, "a_support": h_a_support, "a_admins": h_a_admins, "a_pay": h_a_pay,
    "a_resell": h_a_resell, "a_tutorial": h_a_tutorial, "a_emoji": h_a_emoji, "a_stats": h_a_stats,
    "a_channel": h_a_channel,
})


# ======================================================================
#                                 MAIN
# ======================================================================
async def resume_pending():
    """Restart ke baad pending deposits ka auto-check dobara chalu"""
    rows = await fetchall("SELECT id,created FROM deposits WHERE status='pending'")
    for d in rows:
        if d["created"] < now() - POLL_MINUTES * 60:
            await run("UPDATE deposits SET status='expired', decided=? WHERE id=?", (now(), d["id"]))
        else:
            spawn(poll_deposit(d["id"]))


async def main():
    global bot, BOT_USERNAME
    logging.basicConfig(level=logging.INFO)
    await init_db()
    bot = Bot(BOT_TOKEN, default=DefaultBotProperties(parse_mode=ParseMode.HTML))
    BOT_USERNAME = (await bot.get_me()).username
    bot.session.middleware(PremiumText())
    dp = Dispatcher(storage=MemoryStorage())
    dp.message.outer_middleware(Guard())
    dp.callback_query.outer_middleware(Guard())
    dp.include_router(r)
    await bot.delete_webhook(drop_pending_updates=True)
    await resume_pending()
    spawn(backup_loop())
    logging.info("Bot started as @%s | DB: %s", BOT_USERNAME, DB_PATH)
    await dp.start_polling(bot)


if __name__ == "__main__":
    asyncio.run(main())
